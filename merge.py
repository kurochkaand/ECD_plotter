"""Merge Windows instrument CSV exports and normalize CD and absorbance spectra."""

from __future__ import annotations

import argparse
import csv
import re
import sys
from datetime import datetime
from io import StringIO
from itertools import islice
from math import isfinite
from os.path import commonprefix
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from scipy.signal import savgol_coeffs


DEFAULT_FOLDER = Path(
    r"X:\Cations_and_PGA\ECD-Szeged\K.Andrii\260730-PLGA-titration-Zn"
)
SPECTRAL_COLUMNS = ["wavelength_nm", "cd_mdeg", "ht_v", "absorbance"]
CHANNEL_SHEETS = {"CD": "cd_mdeg", "HT": "ht_v", "Abs": "absorbance"}
CONCENTRATION_COLUMNS = ["spectrum", "concentration_coefficient", "Ph", "short name"]


def read_acquisition_datetime(file: Path) -> datetime:
    """Read YY/MM/DD and a 24-hour or AM/PM time from cells B5 and B6."""
    with file.open(encoding="latin-1", newline="") as stream:
        rows = list(islice(csv.reader(stream), 6))
    date_text = rows[4][1].strip() if len(rows) > 4 and len(rows[4]) > 1 else ""
    time_text = rows[5][1].strip() if len(rows) > 5 and len(rows[5]) > 1 else ""

    for date_format in ("%y/%m/%d", "%Y/%m/%d"):
        for time_format in ("%H:%M:%S", "%I:%M:%S %p"):
            try:
                return datetime.strptime(
                    f"{date_text} {time_text}", f"{date_format} {time_format}"
                )
            except ValueError:
                continue
    raise ValueError(
        f"Invalid acquisition date/time in {file}\n"
        f"  B5 (CSV line 5): {date_text!r}; expected YY/MM/DD or YYYY/MM/DD\n"
        f"  B6 (CSV line 6): {time_text!r}; expected HH:MM:SS or H:MM:SS AM/PM"
    )


def find_spectral_files(folder: Path) -> list[Path]:
    """Find CSV files ordered by acquisition date/time, earliest first."""
    if not folder.is_dir():
        raise NotADirectoryError(f"Data folder not found: {folder}")

    files = sorted(
        (file for file in folder.glob("*.csv") if file.is_file()),
        key=lambda file: (read_acquisition_datetime(file), file.name.casefold()),
    )
    if not files:
        raise FileNotFoundError(f"No CSV files found directly inside {folder}")

    return files


def read_spectrum(file: Path) -> pd.DataFrame:
    """Read the block between XYDATA and Extended Information."""
    lines = file.read_text(encoding="latin-1").splitlines(keepends=True)
    data_start = None
    data_end = None
    for row_number, line in enumerate(lines):
        if line.strip() == "XYDATA":
            data_start = row_number + 1
        elif line.startswith("##### Extended Information"):
            data_end = row_number
            break

    if data_start is None:
        raise ValueError(f"XYDATA marker not found in {file}")
    if data_end is None:
        raise ValueError(f"Extended Information marker not found in {file}")

    data_block = "".join(lines[data_start:data_end])
    if not data_block.strip():
        raise ValueError(f"No spectral data found in {file}")

    frame = pd.read_csv(
        StringIO(data_block),
        header=None,
        index_col=False,
        skip_blank_lines=True,
    )
    if frame.shape[1] < 3:
        raise ValueError(
            f"Expected wavelength, CD, and HT columns with optional Abs in {file}; "
            f"found {frame.shape[1]} columns"
        )
    # Older three-column exports have no absorbance; keep their Abs cells blank.
    frame = frame.reindex(columns=range(len(SPECTRAL_COLUMNS)))
    frame.columns = SPECTRAL_COLUMNS
    frame.attrs["csv_lines"] = [
        row_number + 1
        for row_number in range(data_start, data_end)
        if lines[row_number].strip()
    ]
    frame.attrs["data_end_line"] = data_end + 1
    return frame


def validate_wavelength_axis(
    reference: pd.DataFrame,
    actual: pd.DataFrame,
    reference_file: Path,
    actual_file: Path,
) -> None:
    """Report the first differing wavelength using physical CSV line numbers."""
    expected_axis = reference["wavelength_nm"].tolist()
    actual_axis = actual["wavelength_nm"].tolist()
    for row in range(max(len(expected_axis), len(actual_axis))):
        has_expected = row < len(expected_axis)
        has_actual = row < len(actual_axis)
        expected_value = expected_axis[row] if has_expected else None
        actual_value = actual_axis[row] if has_actual else None
        if has_expected and has_actual and expected_value == actual_value:
            continue

        expected_line = (
            reference.attrs["csv_lines"][row]
            if has_expected else reference.attrs["data_end_line"]
        )
        actual_line = (
            actual.attrs["csv_lines"][row]
            if has_actual else actual.attrs["data_end_line"]
        )
        expected_text = (
            f"{expected_value!r} nm" if has_expected else "<end of spectral data>"
        )
        actual_text = f"{actual_value!r} nm" if has_actual else "<end of spectral data>"
        raise ValueError(
            f"X axis mismatch in {actual_file}\n"
            f"  Reference file: {reference_file}\n"
            f"  Spectral data row: {row + 1}\n"
            f"  Expected: {expected_text} (CSV line {expected_line})\n"
            f"  Actual:   {actual_text} (CSV line {actual_line})\n"
            f"  Data rows: expected {len(expected_axis)}, actual {len(actual_axis)}"
        )


def extract_ph_from_name(spectrum_name: str) -> float | None:
    """Read a pH number from a filename stem, allowing a space after pH."""
    match = re.search(
        r"(?:^|[_\s-])ph\s*([+-]?(?:\d+(?:\.\d+)?|\.\d+))(?=$|[_\s-])",
        spectrum_name,
        flags=re.IGNORECASE,
    )
    return float(match.group(1)) if match else None


def determine_short_names(
    sample_names: list[str], subtraction_reference: str | None = None,
) -> dict[str, str]:
    """Remove shared filename affixes, keeping nonempty, unique spectrum labels."""
    full_names = dict(zip(sample_names, sample_names))
    samples = [name for name in sample_names if name != subtraction_reference]
    if len(samples) < 2:
        return full_names

    prefix = commonprefix(samples)
    remainders = [name[len(prefix):] for name in samples]
    suffix = commonprefix([name[::-1] for name in remainders])[::-1]
    shortened = {
        name: remainder[:len(remainder) - len(suffix)].strip(" _-\t")
        for name, remainder in zip(samples, remainders)
    }
    labels = full_names | shortened
    # Keep original stems if trimming removes a whole name or creates ambiguity.
    # The reference retains its full stem and never participates in affix detection.
    if (
        not all(labels.values())
        or len(set(labels.values())) != len(labels)
        or any(label in {"zero", "wavelength_nm"} for label in shortened.values())
    ):
        return full_names
    return labels


def load_concentration_table(
    folder: Path, sample_names: list[str], *, refresh_ph: bool = False,
    subtraction_reference: str | None = None,
) -> pd.DataFrame:
    """Create or validate coefficients and pH, and regenerate short names."""
    concentration_file = folder / "concentration_table.txt"
    created = not concentration_file.exists()
    if created:
        concentration_table = pd.DataFrame(
            {
                "spectrum": sample_names,
                "concentration_coefficient": 1.0,
            }
        )
    else:
        with concentration_file.open(encoding="utf-8-sig", newline="") as stream:
            header = next(csv.reader(stream, delimiter="\t"), [])
            # Accept extra alignment tabs in a manually edited header.
            column_count = sum(bool(column.strip()) for column in header)
            if column_count not in (2, 3, 4):
                raise ValueError(
                    f"{concentration_file} must have two, three, or four tab-separated columns: "
                    "spectrum name, concentration coefficient, optional Ph and short name"
                )
            rows = []
            for fields in csv.reader(stream, delimiter="\t"):
                fields = [field.strip() for field in fields]
                if not any(fields):
                    continue
                while fields and not fields[-1]:
                    fields.pop()
                # Legacy tables use repeated tabs for alignment. Four-column
                # rows also allow a blank Ph cell before the short name.
                if column_count < 4 or len(fields) > column_count or (len(fields) > 1 and not fields[1]):
                    fields = [field for field in fields if field]
                if len(fields) > column_count:
                    raise ValueError(
                        f"{concentration_file} must have {column_count} tab-separated columns "
                        "per row; found extra values"
                    )
                rows.append([field or None for field in fields] + [None] * (column_count - len(fields)))
            concentration_table = pd.DataFrame(rows, columns=CONCENTRATION_COLUMNS[:column_count])
        concentration_table["spectrum"] = (
            concentration_table["spectrum"].astype(str).str.strip()
        )
        if concentration_table["spectrum"].duplicated().any():
            duplicates = concentration_table.loc[
                concentration_table["spectrum"].duplicated(), "spectrum"
            ].tolist()
            raise ValueError(
                f"Duplicate spectra in {concentration_file}: {', '.join(duplicates)}"
            )

    inferred_ph = concentration_table["spectrum"].map(extract_ph_from_name).astype(float)
    if "Ph" not in concentration_table:
        concentration_table["Ph"] = inferred_ph.fillna(7.0)
    elif refresh_ph:
        concentration_table["Ph"] = inferred_ph.fillna(
            concentration_table["Ph"]
        ).fillna(7.0)
    else:
        concentration_table["Ph"] = concentration_table["Ph"].fillna(
            inferred_ph
        ).fillna(7.0)

    concentration_table["concentration_coefficient"] = pd.to_numeric(
        concentration_table["concentration_coefficient"], errors="coerce"
    )
    coefficients = concentration_table.set_index("spectrum")[
        "concentration_coefficient"
    ]
    missing_spectra = [name for name in sample_names if name not in coefficients.index]
    extra_spectra = [name for name in coefficients.index if name not in sample_names]
    if missing_spectra or extra_spectra:
        problems = []
        if missing_spectra:
            problems.append(f"missing: {', '.join(missing_spectra)}")
        if extra_spectra:
            problems.append(f"not present in data folder: {', '.join(extra_spectra)}")
        raise ValueError(
            f"{concentration_file} does not match the CSV files "
            f"({'; '.join(problems)})"
        )

    invalid = coefficients[
        ~coefficients.map(isfinite) | (coefficients <= 0)
    ].index.tolist()
    if invalid:
        raise ValueError(
            "Concentration coefficients must be finite and greater than zero for: "
            + ", ".join(invalid)
        )
    concentration_table = concentration_table.set_index("spectrum").loc[sample_names]
    concentration_table["short name"] = pd.Series(
        determine_short_names(sample_names, subtraction_reference)
    )
    concentration_table.to_csv(concentration_file, sep="\t", index=True)
    if created:
        print(f"Created {concentration_file}")
    return concentration_table


def load_concentration_coefficients(
    folder: Path, sample_names: list[str], *, refresh_ph: bool = False,
    subtraction_reference: str | None = None,
) -> pd.Series:
    """Return coefficients by filename while updating the concentration table."""
    return load_concentration_table(
        folder, sample_names, refresh_ph=refresh_ph,
        subtraction_reference=subtraction_reference,
    )["concentration_coefficient"]


def select_subtraction_reference(sample_names: list[str], prefix: str | None) -> str | None:
    """Require one matching spectrum only when subtraction is requested."""
    if prefix is None:
        return None
    normalized_prefix = prefix.strip().casefold()
    if not normalized_prefix:
        raise ValueError("The --subtract filename prefix cannot be empty.")
    matches = [name for name in sample_names if name.casefold().startswith(normalized_prefix)]
    if len(matches) != 1:
        details = f": {', '.join(matches)}" if matches else ""
        raise ValueError(
            f"Expected exactly one spectrum starting with {prefix!r}; "
            f"found {len(matches)}{details}"
        )
    return matches[0]


def corrected_spectral_columns(
    spectral_data: dict[str, pd.Series],
    coefficients: pd.Series,
    subtraction_reference: str | None,
) -> list[pd.Series]:
    """Normalize a channel, using Excel formulas when subtracting a reference."""
    if subtraction_reference is None:
        return [
            (signal / coefficients[name]).rename(name)
            for name, signal in spectral_data.items()
        ]

    reference_column = get_column_letter(list(spectral_data).index(subtraction_reference) + 2)
    corrected_columns = []
    for column, (name, signal) in enumerate(spectral_data.items(), start=2):
        if name == subtraction_reference:
            continue
        source_column = get_column_letter(column)
        coefficient = float(coefficients[name])
        formulas = []
        for position, row in enumerate(range(2, len(signal) + 2)):
            if pd.isna(signal.iloc[position]) or pd.isna(spectral_data[subtraction_reference].iloc[position]):
                formulas.append(None)
                continue
            difference = f"{source_column}{row}-${reference_column}{row}"
            formulas.append(
                f"={difference}" if coefficient == 1.0
                else f"=({difference})/{coefficient!r}"
            )
        corrected_columns.append(pd.Series(formulas, index=signal.index, name=name))
    return corrected_columns


def smooth_cd_spectra(
    cd_data: dict[str, pd.Series], window_points: int, *, allow_gaps: bool = False
) -> dict[str, pd.Series]:
    """Apply quadratic Savitzky-Golay fits without padding or shifting samples."""
    if window_points < 3:
        raise ValueError("The --smooth window must contain at least 3 points for polynomial order 2.")
    for name, signal in cd_data.items():
        if window_points > len(signal):
            raise ValueError(
                f"The --smooth window of {window_points} points exceeds the "
                f"{len(signal)} data points in {name}.csv."
            )

    left = window_points // 2
    right = window_points - left - 1
    weights = savgol_coeffs(window_points, polyorder=2, pos=left, use="dot")
    window_positions = np.arange(window_points, dtype=float)
    smoothed_data = {}
    for name, signal in cd_data.items():
        try:
            values = signal.to_numpy(dtype=float)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"CD values in {name}.csv must be numeric for smoothing.") from exc
        if np.isinf(values).any() or (not allow_gaps and np.isnan(values).any()):
            raise ValueError(f"CD values in {name}.csv must be finite for smoothing.")

        smoothed = values.copy()
        # Never bridge excluded points. Keep short valid runs unsmoothed.
        boundaries = np.flatnonzero(np.diff(np.r_[False, ~np.isnan(values), False]))
        for start, stop in zip(boundaries[::2], boundaries[1::2]):
            if stop - start < window_points:
                continue
            segment = values[start:stop]
            smoothed[start + left:stop - right] = np.convolve(segment, weights[::-1], mode="valid")
            # Fit only measured points at each valid segment's edges.
            first_fit = np.polynomial.Polynomial.fit(window_positions, segment[:window_points], 2)
            last_fit = np.polynomial.Polynomial.fit(window_positions, segment[-window_points:], 2)
            smoothed[start:start + left] = first_fit(window_positions[:left])
            smoothed[stop - right:stop] = last_fit(window_positions[-right:])
        smoothed_data[name] = pd.Series(smoothed, index=signal.index, name=signal.name)
    return smoothed_data


def plot_corrected_spectra(
    wavelengths: pd.Series,
    spectral_channels: dict[str, dict[str, pd.Series]],
    coefficients: pd.Series,
    subtraction_reference: str | None,
    *, smooth_window: int | None = None, ht_threshold: float | None = None,
    short_names: dict[str, str] | None = None,
) -> None:
    """Display corrected CD and absorbance panels without reading Excel formulas."""
    import matplotlib.pyplot as plt

    sample_names = [name for name in spectral_channels["CD"] if name != subtraction_reference]
    if not sample_names:
        raise ValueError("No sample spectra remain to plot after reference subtraction.")

    corrected_channels = {}
    for channel_name, spectral_data in spectral_channels.items():
        corrected = {}
        for name in sample_names:
            signal = spectral_data[name]
            if subtraction_reference is not None:
                signal = signal - spectral_data[subtraction_reference]
            corrected[name] = signal / coefficients[name]
        # Older CSV exports have no absorbance. Only plot available Abs curves.
        if channel_name == "Abs":
            corrected = {name: signal for name, signal in corrected.items() if signal.notna().any()}
            if not corrected:
                continue
        corrected_channels[channel_name] = corrected

    figure, axes = plt.subplots(
        len(corrected_channels), 1, figsize=(14, 6 * len(corrected_channels)),
        sharex=True, squeeze=False,
    )
    try:
        colors = plt.get_cmap("turbo")
        for axis, (channel_name, corrected) in zip(axes[:, 0], corrected_channels.items()):
            for index, name in enumerate(sample_names):
                if name not in corrected:
                    continue
                axis.plot(
                    wavelengths,
                    corrected[name],
                    label=short_names[name] if short_names is not None else name,
                    linewidth=1.3,
                    color=colors(index / max(len(sample_names) - 1, 1)),
                )

            label = "CD" if channel_name == "CD" else "absorbance"
            title = f"Concentration-normalized {label} spectra"
            if subtraction_reference is not None:
                title += f"\nReference subtracted: {subtraction_reference}"
            if channel_name == "CD":
                if smooth_window is not None:
                    title += f"\nSavitzky-Golay: {smooth_window} points, polynomial order 2"
                if ht_threshold is not None:
                    title += f"\nHT cutoff: {ht_threshold:g} V"
            axis.set_title(title)
            axis.set_xlabel("Wavelength (nm)")
            axis.set_ylabel(f"Normalized {label} ({'mdeg' if channel_name == 'CD' else 'AU'})")
            axis.grid(alpha=0.25)
            axis.legend(loc="upper left", bbox_to_anchor=(1.02, 1), fontsize=7)
        figure.tight_layout()
        plt.show(block=True)
    finally:
        plt.close(figure)


def merge_spectra(
    folder: Path, *, refresh_ph: bool = False, subtract: str | None = None,
    plot: bool = False, smooth: int | None = None, ht_threshold: float | None = None,
) -> Path:
    """Write measured channels and corrected CD/Abs sheets on the wavelength axis."""
    if ht_threshold is not None and (not isfinite(ht_threshold) or ht_threshold < 0):
        raise ValueError("The --ht-threshold value must be finite and greater than or equal to zero.")
    files = find_spectral_files(folder)
    reference_frame: pd.DataFrame | None = None
    channel_columns: dict[str, list[pd.Series]] = {
        sheet_name: [] for sheet_name in CHANNEL_SHEETS
    }
    cd_data: dict[str, pd.Series] = {}
    abs_data: dict[str, pd.Series] = {}

    for file in files:
        frame = read_spectrum(file)
        if reference_frame is None:
            reference_frame = frame
            for columns in channel_columns.values():
                columns.append(frame["wavelength_nm"])
        else:
            validate_wavelength_axis(reference_frame, frame, files[0], file)

        if ht_threshold is not None:
            try:
                ht_values = pd.to_numeric(frame["ht_v"], errors="raise")
            except (TypeError, ValueError) as exc:
                raise ValueError(f"HT values in {file} must be numeric for --ht-threshold.") from exc
            if not np.isfinite(ht_values).all():
                raise ValueError(f"HT values in {file} must be finite for --ht-threshold.")
            frame["cd_mdeg"] = frame["cd_mdeg"].mask(ht_values > ht_threshold)

        for sheet_name, column_name in CHANNEL_SHEETS.items():
            channel_columns[sheet_name].append(frame[column_name].rename(file.stem))
        cd_data[file.stem] = frame["cd_mdeg"]
        abs_data[file.stem] = frame["absorbance"]

    sample_names = list(cd_data)
    subtraction_reference = select_subtraction_reference(sample_names, subtract)
    if smooth is not None:
        cd_data = smooth_cd_spectra(cd_data, smooth, allow_gaps=ht_threshold is not None)
    concentration_table = load_concentration_table(
        folder, sample_names, refresh_ph=refresh_ph,
        subtraction_reference=subtraction_reference,
    )
    coefficients = concentration_table["concentration_coefficient"]
    short_names = concentration_table["short name"].to_dict()

    # CD inputs include optional smoothing; Abs inputs retain the measured values.
    spectral_channels = {"CD": cd_data, "Abs": abs_data}
    summaries = {}
    for channel_name, spectral_data in spectral_channels.items():
        wavelengths = channel_columns[channel_name][0]
        summaries[channel_name] = pd.concat(
            [
                wavelengths,
                *(spectral_data[name].rename(name) for name in sample_names),
                pd.Series(0, index=wavelengths.index, name="zero"),
                *corrected_spectral_columns(spectral_data, coefficients, subtraction_reference),
            ],
            axis=1,
        )

    output_file = folder.parent / f"{folder.name}-combined.xlsx"
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        for sheet_name, columns in channel_columns.items():
            pd.concat(columns, axis=1).rename(columns=short_names).to_excel(
                writer, sheet_name=sheet_name, index=False
            )
        for channel_name, summary in summaries.items():
            summary.rename(columns=short_names).to_excel(
                writer, sheet_name=f"{channel_name}_subtr", index=False
            )
        if subtraction_reference is not None:
            writer.book.calculation = CalcProperties(
                calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True
            )
    if plot:
        plot_corrected_spectra(
            channel_columns["CD"][0], spectral_channels, coefficients, subtraction_reference,
            smooth_window=smooth, ht_threshold=ht_threshold, short_names=short_names,
        )
    return output_file


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Merge instrument CSV files and normalize CD and absorbance spectra on Windows."
    )
    parser.add_argument(
        "--folder",
        type=Path,
        default=DEFAULT_FOLDER,
        help="Data-folder override for DEFAULT_FOLDER configured in the script.",
    )
    parser.add_argument(
        "--refresh-ph",
        action="store_true",
        help="Update existing Ph values from filenames, keeping concentration coefficients.",
    )
    parser.add_argument(
        "--subtract",
        metavar="PREFIX",
        help="Subtract the CD and absorbance reference matching this filename prefix using Excel formulas.",
    )
    parser.add_argument(
        "--plot",
        action="store_true",
        help="Display corrected CD and absorbance spectra in separate panels after saving.",
    )
    parser.add_argument(
        "--smooth",
        type=int,
        nargs="?",
        const=50,
        default=None,
        metavar="POINTS",
        help="Smooth CD before subtraction using Savitzky-Golay (order 2, no padding). "
             "Defaults to 50 points when supplied without a value.",
    )
    parser.add_argument(
        "--ht-threshold",
        type=float,
        nargs="?",
        const=500.0,
        default=None,
        metavar="VALUE",
        help="Remove CD points where that spectrum's HT exceeds VALUE (V), before smoothing. "
             "Defaults to 500 when supplied without a value.",
    )
    args = parser.parse_args()

    try:
        output_file = merge_spectra(
            args.folder.resolve(), refresh_ph=args.refresh_ph, subtract=args.subtract,
            plot=args.plot, smooth=args.smooth, ht_threshold=args.ht_threshold,
        )
    except (OSError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    print(f"Saved {output_file}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
