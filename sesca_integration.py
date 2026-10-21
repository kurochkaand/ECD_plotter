"""Run the official SESCA_deconv engine on corrected MRE spectra."""

from __future__ import annotations

import importlib.util
import io
import os
import random
import tempfile
import warnings
from contextlib import redirect_stdout
from copy import deepcopy
from dataclasses import dataclass, field
from hashlib import sha256
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from scipy.optimize import Bounds, LinearConstraint, OptimizeWarning, minimize


SESCA_URL = "https://www.mpinat.mpg.de/sesca"
DEFAULT_SESCA_DIR = Path(__file__).resolve().parent / ".sesca" / "SESCA_v097"
RESULTS_SHEET = "SESCA_results"
FIT_SHEET = "SESCA_fit"


@dataclass(frozen=True)
class SescaOptions:
    directory: Path | None = None
    basis: str = "DS-dT"
    min_nm: float = 190.0
    max_nm: float = 250.0
    repeats: int = 10
    seed: int = 0
    shared_scaling: bool = False
    fixed_scaling: float | None = None


@dataclass
class SpectrumFit:
    name: str
    status: str
    note: str = ""
    fractions: list[float] = field(default_factory=list)
    wavelengths: list[float] = field(default_factory=list)
    observed_mre: list[float] = field(default_factory=list)
    fitted_mre: list[float] = field(default_factory=list)
    scale: float | None = None
    rmsd_kmre: float | None = None
    nrmsd: float | None = None


@dataclass
class SescaAnalysis:
    classes: list[str]
    fits: list[SpectrumFit]
    parameters: list[tuple[str, str | int | float]]


def resolve_sesca(options: SescaOptions) -> tuple[Path, Path]:
    """Find an explicit installation, SESCA_HOME, or the local distribution."""
    if not np.isfinite([options.min_nm, options.max_nm]).all() or not 0 < options.min_nm < options.max_nm:
        raise ValueError("SESCA wavelength bounds must be finite, positive, and increasing.")
    if isinstance(options.repeats, bool) or not isinstance(options.repeats, int) or options.repeats < 1:
        raise ValueError("SESCA repeats must be a positive integer.")
    if isinstance(options.seed, bool) or not isinstance(options.seed, int) or not 0 <= options.seed < 2**32:
        raise ValueError("SESCA seed must be an integer between 0 and 4294967295.")
    if options.fixed_scaling is not None and (
        isinstance(options.fixed_scaling, bool) or not np.isfinite(options.fixed_scaling) or options.fixed_scaling <= 0
    ):
        raise ValueError("The fixed shared scaling factor must be finite and greater than zero.")
    directory = Path(options.directory or os.environ.get("SESCA_HOME") or DEFAULT_SESCA_DIR).resolve()
    if not (directory / "scripts" / "SESCA_deconv.py").is_file():
        raise FileNotFoundError(
            f"SESCA_deconv.py not found in {directory / 'scripts'}. "
            f"Download and extract SESCA from {SESCA_URL} to {DEFAULT_SESCA_DIR}, "
            "or pass --sesca-dir PATH to its root directory. "
            "MRE conversion is performed by merge.py and does not require SESCA."
        )
    basis = Path(options.basis)
    if not basis.is_file():
        basis = directory / "libs" / f"Map_{options.basis}.dat"
    if not basis.is_file():
        raise FileNotFoundError(
            f"SESCA basis set not found: {options.basis!r}. "
            "Use a standalone basis name such as DS-dT, DS5-4, or DSSP-1, "
            "or pass the path to a SESCA basis file with --sesca-basis."
        )
    return directory, basis.resolve()


def read_basis(basis_file: Path, options: SescaOptions) -> tuple[list[str], np.ndarray]:
    """Read SESCA's labeled basis table, retaining its native wavelength grid."""
    classes = []
    rows = []
    for line in basis_file.read_text(encoding="ascii").splitlines():
        if line.startswith("#|"):
            classes = line[2:].split()[1:]
        elif line.strip() and not line.lstrip().startswith("#"):
            try:
                row = [float(value) for value in line.split()]
            except ValueError as exc:
                raise ValueError(f"Invalid numeric basis data in {basis_file}.") from exc
            if row and options.min_nm <= row[0] <= options.max_nm:
                rows.append(row)
    if not classes or not rows or any(len(row) != len(classes) + 1 for row in rows):
        raise ValueError(f"No valid SESCA basis table in the selected wavelength range: {basis_file}")
    values = np.asarray(rows, dtype=float)
    if not np.isfinite(values).all() or len(set(values[:, 0])) != len(values):
        raise ValueError(f"SESCA basis data must be finite with unique wavelengths: {basis_file}")
    return classes, values


def load_engine(directory: Path):
    """Import the official engine without changing the upstream installation."""
    script = directory / "scripts" / "SESCA_deconv.py"
    spec = importlib.util.spec_from_file_location("_ecd_sesca_deconv", script)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load SESCA from {script}.")
    module = importlib.util.module_from_spec(spec)
    try:
        with redirect_stdout(io.StringIO()):
            spec.loader.exec_module(module)
    except (ImportError, SyntaxError) as exc:
        raise RuntimeError(f"Cannot load SESCA: {exc}") from exc
    return module


def refine_shared_scaling(
    inputs: list[tuple[SpectrumFit, np.ndarray]], fixed_scale: float | None = None,
) -> list[tuple[str, str | int | float]]:
    """Jointly minimize unscaled spectral SSE with equal sums of basis weights.

    SESCA fits observed = basis @ weights, then reports fractions = weights / q
    and scaling_factor = 1 / q, where q is the sum of weights. Here every spectrum
    has its own nonnegative weights but shares q. This is a convex least-squares
    problem with linear equality constraints, not an average of individual fits.
    A supplied scaling factor fixes q = 1 / fixed_scale during optimization.
    """
    if not inputs:
        parameters = [
            ("sesca_shared_scaling_source", "fixed" if fixed_scale is not None else "automatic"),
            ("sesca_shared_fit_status", "No eligible spectra"),
        ]
        if fixed_scale is not None:
            parameters.append(("sesca_shared_scaling_factor", fixed_scale))
        return parameters
    count, classes = len(inputs), inputs[0][1].shape[1]
    points = sum(len(fit.wavelengths) for fit, _ in inputs)
    dimensions = 1 + count * classes
    design = np.zeros((points, dimensions))
    target = np.empty(points)
    equality = np.zeros((count, dimensions))
    initial = np.empty(dimensions)
    initial[0] = 1.0
    # Work near unit-sized weights even when measured intensity is misnormalized.
    # This reference only conditions the problem; the final factor is optimized.
    reference_sum = 1.0 / (fixed_scale if fixed_scale is not None else float(np.median([fit.scale for fit, _ in inputs])))
    if not np.isfinite(reference_sum) or reference_sum <= 0:
        raise RuntimeError("Cannot initialize a finite positive shared coefficient sum.")
    offset = 0
    for index, (fit, matrix) in enumerate(inputs):
        columns = slice(1 + index * classes, 1 + (index + 1) * classes)
        rows = slice(offset, offset + len(fit.wavelengths))
        design[rows, columns] = matrix
        target[rows] = np.asarray(fit.observed_mre) / 1000.0 / reference_sum
        equality[index, 0] = -1.0
        equality[index, columns] = 1.0
        initial[columns] = fit.fractions
        offset = rows.stop
    magnitude = max(float(np.max(np.abs(design))), float(np.max(np.abs(target))))
    design /= magnitude
    target /= magnitude

    def objective(values):
        residual = design @ values - target
        return float(residual @ residual)

    def gradient(values):
        return 2.0 * design.T @ (design @ values - target)

    lower, upper = np.zeros(dimensions), np.full(dimensions, np.inf)
    if fixed_scale is not None:
        lower[0] = upper[0] = 1.0
    result = minimize(
        objective, initial, jac=gradient, method="SLSQP",
        bounds=Bounds(lower, upper),
        constraints=LinearConstraint(equality, 0.0, 0.0),
        options={"ftol": 1e-14, "maxiter": 5000},
    )
    if not result.success:
        raise RuntimeError(f"Shared-scaling fit failed: {result.message}")
    if (not np.isfinite(result.x).all() or np.any(result.x < -1e-9)
            or not np.allclose(equality @ result.x, 0.0, atol=1e-8)
            or result.x[0] <= 1e-12):
        raise RuntimeError("Shared-scaling fit did not produce valid positive scaling and nonnegative weights.")
    shared_scale = float(fixed_scale) if fixed_scale is not None else 1.0 / (reference_sum * float(result.x[0]))
    if not np.isfinite(shared_scale) or shared_scale <= 0:
        raise RuntimeError("Shared-scaling fit did not produce a finite positive scaling factor.")

    unscaled_sse = 0.0
    for index, (fit, matrix) in enumerate(inputs):
        weights = result.x[1 + index * classes:1 + (index + 1) * classes]
        fractions = np.maximum(weights, 0.0)
        total = float(fractions.sum())
        if total <= 0 or not np.isclose(total, result.x[0], rtol=1e-7, atol=1e-10):
            raise RuntimeError("Shared-scaling fit returned inconsistent coefficient sums.")
        fractions /= total  # Remove only numerical equality/bound roundoff.
        predicted = matrix @ fractions
        observed = np.asarray(fit.observed_mre) / 1000.0
        scaled = shared_scale * observed
        residual = scaled - predicted
        fit.fractions = fractions.tolist()
        fit.fitted_mre = (predicted * 1000.0).tolist()
        fit.scale = shared_scale
        fit.rmsd_kmre = float(np.sqrt(np.mean(residual**2)))
        fit.nrmsd = fit.rmsd_kmre / float(np.sqrt(np.mean(scaled**2)))
        unscaled_sse += float(np.sum((observed - predicted / shared_scale)**2))
    return [
        ("sesca_shared_scaling_factor", shared_scale),
        ("sesca_shared_scaling_source", "fixed" if fixed_scale is not None else "automatic"),
        ("sesca_shared_fit_status", "Converged"),
        ("sesca_shared_fit_method", "Joint constrained least squares (SLSQP)"),
        ("sesca_shared_fit_iterations", int(result.nit)),
        ("sesca_shared_fit_objective", "Sum of squared (observed_kMRE - fitted_kMRE / scaling_factor) residuals"),
        ("sesca_shared_fit_RMSD_kMRE_unscaled", float(np.sqrt(unscaled_sse / points))),
    ]


def fit_spectra(
    names: list[str], mre_rows: list[list[float | None]], options: SescaOptions
) -> SescaAnalysis:
    """Fit valid basis-grid points only; never interpolate over excluded CD data."""
    directory, basis_file = resolve_sesca(options)
    classes, basis = read_basis(basis_file, options)
    engine = load_engine(directory)
    defaults = deepcopy(engine.Pass_Defaults())
    wavelengths = np.asarray([row[0] for row in mre_rows], dtype=float)
    if not np.isfinite(wavelengths).all() or len(set(wavelengths)) != len(wavelengths):
        raise ValueError("SESCA input must have finite, unique wavelengths.")
    matches = []
    for basis_row in basis:
        positions = np.flatnonzero(np.isclose(wavelengths, basis_row[0], rtol=0, atol=1e-6))
        if len(positions) > 1:
            raise ValueError(f"Multiple input wavelengths match SESCA basis wavelength {basis_row[0]:g} nm.")
        if len(positions):
            matches.append((int(positions[0]), basis_row))

    fits = []
    shared_inputs = []
    with tempfile.TemporaryDirectory(prefix="ECD SESCA ") as temporary:
        spectrum_file = Path(temporary) / "spectrum.dat"
        for column, name in enumerate(names, start=1):
            selected = [(index, row) for index, row in matches if mre_rows[index][column] is not None]
            minimum_points = max(10, 2 * len(classes))
            if len(selected) < minimum_points:
                fits.append(SpectrumFit(name, "skipped", f"Only {len(selected)} valid basis-grid points; need {minimum_points}."))
                continue
            matrix = np.asarray([row[1:] for _, row in selected])
            fit_waves = np.asarray([row[0] for _, row in selected])
            observed = np.asarray([mre_rows[index][column] for index, _ in selected], dtype=float)
            if not np.isfinite(observed).all():
                raise ValueError(f"Nonfinite MRE input for SESCA spectrum {name!r}.")
            if np.ptp(fit_waves) < 20 or np.linalg.matrix_rank(matrix) < len(classes):
                fits.append(SpectrumFit(name, "skipped", "Need at least 20 nm coverage and independent basis spectra."))
                continue
            if not np.any(observed):
                fits.append(SpectrumFit(name, "skipped", "All selected MRE values are zero."))
                continue
            # SESCA and its libraries use kMRE, rather than the workbook's MRE.
            np.savetxt(spectrum_file, np.column_stack([fit_waves, observed / 1000.0]), fmt="%.12g")
            log = io.StringIO()
            python_state, numpy_state = random.getstate(), np.random.get_state()
            try:
                random.seed(options.seed)
                np.random.seed(options.seed)
                # Read_Args mutates SESCA's defaults. Reset them for each spectrum.
                engine.Def_Args = deepcopy(defaults)
                with redirect_stdout(log), warnings.catch_warnings():
                    # SESCA 0.97 passes a retired SciPy option; modern defaults apply.
                    warnings.filterwarnings("ignore", message="Unknown solver options: ftol", category=OptimizeWarning)
                    arguments = engine.Read_Args([
                        "@spect", str(spectrum_file), "@lib", str(basis_file),
                        "@write", "0", "@range", f"{options.min_nm},{options.max_nm}",
                        "@rep", str(options.repeats), "@iter", "5000", "@err", "0", "@verb", "0",
                    ])
                    result = engine.Deconv_Main(arguments)
            except (Exception, SystemExit) as exc:
                raise RuntimeError(f"SESCA failed for {name!r}: {exc}\n{log.getvalue()[-1500:]}") from exc
            finally:
                random.setstate(python_state)
                np.random.set_state(numpy_state)
            fractions = np.asarray(result[0], dtype=float)
            rmsd, _, nrmsd, scale = map(float, result[1:5])
            if (fractions.shape != (len(classes),) or not np.isfinite(fractions).all()
                    or np.any(fractions < 0) or not np.isclose(fractions.sum(), 1, atol=1e-6)
                    or not np.isfinite([rmsd, nrmsd, scale]).all() or min(rmsd, nrmsd) < 0 or scale <= 0):
                raise RuntimeError(f"SESCA returned invalid fit values for {name!r}.")
            fitted = matrix @ fractions * 1000.0
            expected_rmsd = float(np.sqrt(np.mean(((observed * scale - fitted) / 1000.0) ** 2)))
            if not np.isclose(rmsd, expected_rmsd, rtol=1e-5, atol=1e-7):
                raise RuntimeError(f"SESCA fitted spectrum and reported RMSD disagree for {name!r}.")
            fit = SpectrumFit(
                name, "ok", fractions=fractions.tolist(), wavelengths=fit_waves.tolist(),
                observed_mre=observed.tolist(), fitted_mre=fitted.tolist(),
                scale=scale, rmsd_kmre=rmsd, nrmsd=nrmsd,
            )
            fits.append(fit)
            shared_inputs.append((fit, matrix))
    shared_scaling = options.shared_scaling or options.fixed_scaling is not None
    shared_parameters = refine_shared_scaling(shared_inputs, options.fixed_scaling) if shared_scaling else []
    return SescaAnalysis(classes, fits, [
        ("secondary_structure_analysis", "SESCA basis with shared scaling" if shared_scaling else "SESCA_deconv"),
        ("sesca_directory", str(directory)),
        ("sesca_source", SESCA_URL),
        ("sesca_basis_file", str(basis_file)),
        ("sesca_basis_sha256", sha256(basis_file.read_bytes()).hexdigest()),
        ("sesca_script_sha256", sha256((directory / "scripts" / "SESCA_deconv.py").read_bytes()).hexdigest()),
        ("sesca_min_nm", options.min_nm), ("sesca_max_nm", options.max_nm),
        ("sesca_repeats", options.repeats), ("sesca_max_iterations", 5000),
        ("sesca_seed", options.seed), ("sesca_input_units", "kMRE = MRE / 1000"),
        ("sesca_rescaling", "Enabled; scaled_observed_MRE = observed_MRE * scaling_factor"),
        ("sesca_scaling_mode", "shared" if shared_scaling else "per-spectrum"),
        ("sesca_fitted_spectra", sum(fit.status == "ok" for fit in fits)),
        ("sesca_skipped_spectra", sum(fit.status == "skipped" for fit in fits)),
        *shared_parameters,
    ])


def write_results(workbook, analysis: SescaAnalysis) -> None:
    """Write fractions and spectra with explicit units and scaling semantics."""
    results = workbook.create_sheet(RESULTS_SHEET)
    results.append([
        "spectrum", "status", "note", *[f"{name}_fraction" for name in analysis.classes],
        "scaling_factor", "RMSD_kMRE", "NRMSD", "fitted_points", "min_nm", "max_nm",
    ])
    curves = workbook.create_sheet(FIT_SHEET)
    curves.append([
        "spectrum", "wavelength_nm", "observed_MRE", "scaled_observed_MRE",
        "fitted_MRE", "residual_MRE",
    ])
    for fit in analysis.fits:
        results.append([
            fit.name, fit.status, fit.note, *(fit.fractions or [None] * len(analysis.classes)),
            fit.scale, fit.rmsd_kmre, fit.nrmsd, len(fit.wavelengths),
            min(fit.wavelengths) if fit.wavelengths else None,
            max(fit.wavelengths) if fit.wavelengths else None,
        ])
        for column in range(4, 4 + len(analysis.classes)):
            results.cell(results.max_row, column).number_format = "0.00%"
        for wave, observed, predicted in zip(fit.wavelengths, fit.observed_mre, fit.fitted_mre):
            scaled = observed * fit.scale
            curves.append([fit.name, wave, observed, scaled, predicted, scaled - predicted])
    for sheet in (results, curves):
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions["A"].width = 65
        sheet.row_dimensions[1].height = 32
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column in range(2, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 22
    results.column_dimensions["B"].width = 12
    results.column_dimensions["C"].width = 75
    for row in results.iter_rows(min_row=2, min_col=4 + len(analysis.classes), max_col=6 + len(analysis.classes)):
        for cell in row:
            cell.number_format = "0.000000"
    for row in curves.iter_rows(min_row=2, min_col=3):
        for cell in row:
            cell.number_format = "0.000000"


def plot_sesca_results(workbook_file: Path) -> None:
    """Display saved structure fractions as stacked percentages in spectrum order."""
    workbook = load_workbook(workbook_file, read_only=True, data_only=True)
    try:
        if RESULTS_SHEET not in workbook.sheetnames:
            raise ValueError(f"No {RESULTS_SHEET} sheet to plot in {workbook_file}.")
        rows = workbook[RESULTS_SHEET].iter_rows(values_only=True)
        headers = list(next(rows))
        class_columns = [i for i, name in enumerate(headers) if isinstance(name, str) and name.endswith("_fraction")]
        if not class_columns or "spectrum" not in headers or "status" not in headers:
            raise ValueError(f"{RESULTS_SHEET} must contain spectrum, status, and structure fraction columns.")
        records = list(rows)
    finally:
        workbook.close()

    name_column, status_column = headers.index("spectrum"), headers.index("status")
    valid = [row[status_column] == "ok" for row in records]
    if not any(valid):
        print(f"No fitted spectra in {RESULTS_SHEET}; no plot to display.")
        return
    fractions = np.asarray([
        [row[column] for column in class_columns] if fitted else [0.0] * len(class_columns)
        for row, fitted in zip(records, valid)
    ], dtype=float)
    if (not np.isfinite(fractions).all() or np.any(fractions < 0) or np.any(fractions > 1)
            or not np.allclose(fractions[valid].sum(axis=1), 1.0, atol=1e-6)):
        raise ValueError(f"Invalid structure fractions in {RESULTS_SHEET}.")

    import matplotlib.pyplot as plt

    figure, axis = plt.subplots(figsize=(16, max(5, 0.45 * len(records) + 2)))
    try:
        positions = np.arange(len(records))
        left = np.zeros(len(records))
        colors = plt.get_cmap("tab20")
        for index, column in enumerate(class_columns):
            percentages = fractions[:, index] * 100.0
            axis.barh(
                positions, percentages, left=left,
                label=headers[column].removesuffix("_fraction"), color=colors(index),
            )
            left += percentages
        for position, fitted in enumerate(valid):
            if not fitted:
                axis.text(2, position, "Skipped", va="center", color="dimgray", fontstyle="italic")
        axis.set_yticks(positions, labels=[row[name_column] for row in records], fontsize=8)
        axis.invert_yaxis()
        axis.set_xlim(0, 100)
        axis.set_xlabel("Secondary structure (%)")
        axis.set_title("SESCA secondary-structure composition")
        axis.set_axisbelow(True)
        axis.grid(axis="x", alpha=0.25)
        axis.legend(loc="upper left", bbox_to_anchor=(1.02, 1))
        figure.tight_layout()
        plt.show(block=True)
    finally:
        plt.close(figure)
