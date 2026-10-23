"""Estimate spectral rank with SVD of the MRE sheet prepared by merge.py."""

from __future__ import annotations

import argparse
import os
import sys
import tempfile
from dataclasses import dataclass
from math import isfinite
from pathlib import Path
from zipfile import BadZipFile

import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from deconvolution import read_mre
from merge import MRE_SHEET, MRE_UNITS
from spectral_component_analysis import (
    CUMULATIVE_VARIANCE_TARGET, SVDResult, summarize_singular_values,
)


RESULTS_SHEET = "MRE_SVD_results"
PARAMETERS_SHEET = "MRE_SVD_parameters"
INPUT_SHEET = "MRE_SVD_input"
PATTERNS_SHEET = "MRE_SVD_patterns"
SCORES_SHEET = "MRE_SVD_scores"
CENTERED_SHEET = "MRE_SVD_centered"
RECONSTRUCTION_SHEET = "MRE_SVD_reconstruction"
COMPARISON_SHEET = "MRE_SVD_comparison"
NOISE_SHEET = "MRE_SVD_noise"
NOISE_WAVELENGTH_SHEET = "MRE_SVD_noise_wavelength"
BASE_SHEETS = (RESULTS_SHEET, PARAMETERS_SHEET, INPUT_SHEET)
DIAGNOSTIC_SHEETS = (
    PATTERNS_SHEET, SCORES_SHEET,
    CENTERED_SHEET, RECONSTRUCTION_SHEET, COMPARISON_SHEET, NOISE_SHEET, NOISE_WAVELENGTH_SHEET,
)
OUTPUT_SHEETS = BASE_SHEETS + DIAGNOSTIC_SHEETS
DEFAULT_MIN_NM = 190.0
DEFAULT_MAX_NM = 250.0


@dataclass(frozen=True)
class SvdDecomposition:
    stats: SVDResult
    scores: np.ndarray
    patterns: np.ndarray
    offset: np.ndarray
    effective_shape: tuple[int, int]


@dataclass(frozen=True)
class Reconstruction:
    rank: int
    reconstructed: np.ndarray
    residuals: np.ndarray
    energy_fraction: float
    relative_error: float
    rms_by_spectrum: np.ndarray
    rms_by_wavelength: np.ndarray
    lag1_correlations: np.ndarray
    lag1_pairs: int


@dataclass(frozen=True)
class MreSvdAnalysis:
    names: list[str]
    wavelengths: np.ndarray
    spectra: np.ndarray  # (spectra, retained wavelength points), in saved MRE units
    svd: SVDResult
    uncentered: SvdDecomposition | None
    centered: SvdDecomposition | None
    min_nm: float
    max_nm: float
    points_in_range: int
    points_outside_range: int
    reconstructions: tuple[Reconstruction, ...]
    diagnostic_components: int
    wavelength_step_nm: float | None

    @property
    def diagnostics_enabled(self) -> bool:
        return self.uncentered is not None

    @property
    def excluded_points(self) -> int:
        return self.points_in_range - len(self.wavelengths)


def decompose_spectra(spectra: np.ndarray, *, centered: bool = False) -> SvdDecomposition:
    """Factor spectra x wavelengths; centered noise has n-1 sample directions."""
    offset = spectra.mean(axis=0) if centered else np.zeros(spectra.shape[1])
    matrix = spectra - offset
    effective_shape = (spectra.shape[0] - int(centered), spectra.shape[1])
    count = min(effective_shape)
    u, singular_values, patterns = np.linalg.svd(matrix, full_matrices=False)
    singular_values, patterns, u = singular_values[:count], patterns[:count], u[:, :count]
    tolerance = np.finfo(float).eps * max(spectra.shape) * np.linalg.norm(spectra)
    if centered and np.linalg.norm(matrix) <= tolerance:
        # Identical spectra have no centered variation; numerical roundoff is not signal.
        singular_values = np.zeros_like(singular_values)
    stats = summarize_singular_values(singular_values, effective_shape, allow_zero_rank=centered)
    scores = u * singular_values
    # Pair the arbitrary sign changes so a component's largest loading is positive.
    signs = np.sign(patterns[np.arange(count), np.argmax(np.abs(patterns), axis=1)])
    signs[signs == 0] = 1
    return SvdDecomposition(
        stats=stats, scores=scores * signs, patterns=patterns * signs[:, None],
        offset=offset, effective_shape=effective_shape,
    )


def reconstruct_spectra(
    spectra: np.ndarray, decomposition: SvdDecomposition, rank: int,
    wavelengths: np.ndarray, wavelength_step_nm: float | None,
) -> Reconstruction:
    """Compare an uncentered rank-k model and inspect its fitted residuals."""
    reconstructed = decomposition.scores[:, :rank] @ decomposition.patterns[:rank] + decomposition.offset
    residuals = spectra - reconstructed
    lag1 = np.full(len(spectra), np.nan)
    neighbors = np.zeros(len(wavelengths) - 1, dtype=bool)
    if wavelength_step_nm is not None:
        neighbors = np.isclose(np.diff(wavelengths), wavelength_step_nm, rtol=1e-6, atol=1e-8)
    pairs = int(neighbors.sum())
    if pairs >= 3:
        for index, residual in enumerate(residuals):
            left, right = residual[:-1][neighbors], residual[1:][neighbors]
            left, right = left - left.mean(), right - right.mean()
            precision = np.finfo(float).eps * max(spectra.shape) * np.linalg.norm(spectra[index])
            left_norm, right_norm = np.linalg.norm(left), np.linalg.norm(right)
            if left_norm > precision and right_norm > precision:
                lag1[index] = np.clip(np.dot(left / left_norm, right / right_norm), -1, 1)
    return Reconstruction(
        rank=rank, reconstructed=reconstructed, residuals=residuals,
        energy_fraction=float(decomposition.stats.cumulative_energy_ratio[rank - 1]),
        relative_error=float(np.linalg.norm(residuals) / np.linalg.norm(spectra)),
        rms_by_spectrum=np.sqrt(np.mean(residuals**2, axis=1)),
        rms_by_wavelength=np.sqrt(np.mean(residuals**2, axis=0)),
        lag1_correlations=lag1, lag1_pairs=pairs,
    )


def analyze_mre(
    sheet: Worksheet, *, min_nm: float = DEFAULT_MIN_NM, max_nm: float = DEFAULT_MAX_NM,
    reconstruction_ranks: list[int] | None = None, diagnostic_components: int = 5,
) -> MreSvdAnalysis:
    """Keep complete shared wavelengths; never interpolate or renormalize MRE."""
    if not isfinite(min_nm) or not isfinite(max_nm) or not 0 < min_nm < max_nm:
        raise ValueError("Wavelength bounds must be finite, positive, and increasing.")
    if isinstance(diagnostic_components, bool) or not isinstance(diagnostic_components, int) or diagnostic_components < 1:
        raise ValueError("Diagnostic component count must be a positive integer.")

    names, rows = read_mre(sheet)
    if len(names) < 2:
        raise ValueError("MRE component analysis requires at least two spectra.")
    values = np.asarray(rows, dtype=float)
    wavelengths = values[:, 0]
    if len(np.unique(wavelengths)) != len(wavelengths):
        raise ValueError("MRE wavelengths must be unique. Rerun merge.py to regenerate the workbook.")

    in_range = (wavelengths >= min_nm) & (wavelengths <= max_nm)
    points_in_range = int(np.count_nonzero(in_range))
    if not points_in_range:
        raise ValueError(
            f"No MRE wavelengths in {min_nm:g}-{max_nm:g} nm; "
            f"the workbook spans {wavelengths.min():g}-{wavelengths.max():g} nm."
        )
    # read_mre rejects formulas, text, errors and infinities; only blanks become NaN.
    complete = in_range & np.isfinite(values[:, 1:]).all(axis=1)
    retained = values[complete]
    if len(retained) < 5:
        raise ValueError(
            f"Only {len(retained)} complete MRE wavelength rows remain in "
            f"{min_nm:g}-{max_nm:g} nm after excluding "
            f"{points_in_range - len(retained)} rows with missing values; need at least 5. "
            "Check HT exclusions or choose a different wavelength range."
        )
    retained = retained[np.argsort(retained[:, 0])]
    spectra = retained[:, 1:].T
    maximum_rank = min(spectra.shape)
    uncentered = centered = None
    reconstructions = ()
    wavelength_step_nm = None
    if reconstruction_ranks is not None:
        if not reconstruction_ranks or any(
            isinstance(rank, bool) or not isinstance(rank, int) or not 1 <= rank <= maximum_rank
            for rank in reconstruction_ranks
        ):
            raise ValueError(f"Reconstruction ranks must be integers from 1 to {maximum_rank}.")
        source_steps = np.diff(np.sort(wavelengths[in_range]))
        step = float(np.median(source_steps))
        wavelength_step_nm = step if np.allclose(source_steps, step, rtol=1e-6, atol=1e-8) else None
        uncentered = decompose_spectra(spectra)
        centered = decompose_spectra(spectra, centered=True)
        svd = uncentered.stats
        reconstructions = tuple(
            reconstruct_spectra(spectra, uncentered, rank, retained[:, 0], wavelength_step_nm)
            for rank in sorted(set(reconstruction_ranks))
        )
    else:
        singular_values = np.linalg.svd(spectra, full_matrices=False, compute_uv=False)
        svd = summarize_singular_values(singular_values, spectra.shape)
    return MreSvdAnalysis(
        names=names, wavelengths=retained[:, 0], spectra=spectra, svd=svd, uncentered=uncentered, centered=centered,
        min_nm=min_nm, max_nm=max_nm, points_in_range=points_in_range,
        points_outside_range=int(len(values) - points_in_range),
        reconstructions=reconstructions,
        diagnostic_components=min(diagnostic_components, maximum_rank) if uncentered is not None else 0,
        wavelength_step_nm=wavelength_step_nm,
    )


def write_results(workbook: Workbook, analysis: MreSvdAnalysis) -> None:
    """Replace only SVD outputs; retain the source spectra and other analyses."""
    for title in OUTPUT_SHEETS:
        if title in workbook.sheetnames:
            del workbook[title]

    results = workbook.create_sheet(RESULTS_SHEET)
    results.append([
        "component", "singular_value_MRE", "spectral_energy_fraction",
        "cumulative_energy_fraction", "above_noise_threshold",
    ])
    svd = analysis.svd
    for index, singular_value in enumerate(svd.singular_values):
        results.append([
            index + 1, float(singular_value), float(svd.explained_energy_ratio[index]),
            float(svd.cumulative_energy_ratio[index]),
            bool(singular_value > svd.gavish_donoho_threshold),
        ])
    for row in results.iter_rows(min_row=2):
        row[1].number_format = "0.000000E+00"
        row[2].number_format = row[3].number_format = "0.00%"

    parameters = workbook.create_sheet(PARAMETERS_SHEET)
    parameters.append(["parameter", "value"])
    for key, value in (
        ("analysis", "Singular value decomposition (SVD)"),
        ("source_sheet", MRE_SHEET),
        ("MRE_units", MRE_UNITS),
        ("matrix_orientation", "Rows: spectra; columns: retained wavelengths"),
        ("mean_centered", False),
        ("normalization", "None; use saved numeric MRE without further conversion"),
        ("spectra_count", len(analysis.names)),
        ("requested_min_nm", analysis.min_nm),
        ("requested_max_nm", analysis.max_nm),
        ("retained_min_nm", float(analysis.wavelengths.min())),
        ("retained_max_nm", float(analysis.wavelengths.max())),
        ("wavelength_rows_in_range", analysis.points_in_range),
        ("retained_wavelength_rows", len(analysis.wavelengths)),
        ("excluded_missing_rows", analysis.excluded_points),
        ("outside_range_rows", analysis.points_outside_range),
        ("missing_value_policy", "Exclude wavelength row if any spectrum has a blank; no interpolation"),
        ("gavish_donoho_rank", svd.gavish_donoho_rank),
        ("gavish_donoho_threshold_MRE", svd.gavish_donoho_threshold),
        ("noise_threshold_method", "Gavish-Donoho unknown-noise threshold; polynomial approximation"),
        ("spectral_energy_target", CUMULATIVE_VARIANCE_TARGET),
        ("spectral_energy_rank", svd.variance_rank),
        ("component_limit", "None; report every singular value"),
        ("diagnostics_enabled", analysis.diagnostics_enabled),
        ("interpretation", "Estimated independent spectral patterns; not a chemical species count"),
    ):
        parameters.append([key, value])
        if key == "gavish_donoho_threshold_MRE":
            parameters.cell(parameters.max_row, 2).number_format = "0.000000E+00"
        elif key == "spectral_energy_target":
            parameters.cell(parameters.max_row, 2).number_format = "0.00%"

    inputs = workbook.create_sheet(INPUT_SHEET)
    inputs.append(["wavelength_nm", *analysis.names])
    for wavelength, signals in zip(analysis.wavelengths, analysis.spectra.T):
        inputs.append([float(wavelength), *[float(value) for value in signals]])
    for row in inputs.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = "0.000000"

    if analysis.diagnostics_enabled:
        write_diagnostics(workbook, analysis)
    output_sheets = OUTPUT_SHEETS if analysis.diagnostics_enabled else BASE_SHEETS
    for title in output_sheets:
        sheet = workbook[title]
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 36
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 28
    inputs.column_dimensions["A"].width = 18
    parameters.column_dimensions["A"].width = 36
    parameters.column_dimensions["B"].width = 90
    for row in parameters.iter_rows(min_row=2, min_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        parameters.row_dimensions[row[0].row].height = 32
    if analysis.diagnostics_enabled:
        workbook[NOISE_SHEET].column_dimensions["G"].width = 65


def write_diagnostics(workbook: Workbook, analysis: MreSvdAnalysis) -> None:
    """Write the optional outputs requested by --reconstruction-ranks."""
    assert analysis.uncentered is not None and analysis.centered is not None
    svd = analysis.svd
    centered_stats = analysis.centered.stats
    centered_results = workbook.create_sheet(CENTERED_SHEET)
    centered_results.append([
        "component", "singular_value_MRE", "variance_fraction",
        "cumulative_variance_fraction", "above_noise_threshold",
    ])
    for index, value in enumerate(centered_stats.singular_values):
        centered_results.append([
            index + 1, float(value), float(centered_stats.explained_energy_ratio[index]),
            float(centered_stats.cumulative_energy_ratio[index]),
            bool(value > centered_stats.gavish_donoho_threshold),
        ])
    for row in centered_results.iter_rows(min_row=2):
        row[1].number_format = "0.000000E+00"
        row[2].number_format = row[3].number_format = "0.00%"
    parameters = workbook[PARAMETERS_SHEET]
    for key, value in (
        ("plotted_components", analysis.diagnostic_components),
        ("reconstruction_ranks", ", ".join(str(model.rank) for model in analysis.reconstructions)),
        ("pattern_definition", "Unit-L2 right singular vectors (Vt); signed mathematical wavelength patterns"),
        ("score_definition", "U times singular values; signed MRE coefficients, not composition fractions"),
        ("centered_definition", "Subtract the mean across spectra at each wavelength; add offset_MRE back to reconstruct"),
        ("centered_effective_samples", analysis.centered.effective_shape[0]),
        ("centered_gavish_donoho_rank", centered_stats.gavish_donoho_rank),
        ("centered_gavish_donoho_threshold_MRE", centered_stats.gavish_donoho_threshold),
        ("centered_variance_rank", centered_stats.variance_rank),
        ("centered_variation_fraction_of_total", float(
            np.linalg.norm(analysis.centered.scores)**2 / np.linalg.norm(analysis.spectra)**2
        )),
        ("lag1_wavelength_step_nm", analysis.wavelength_step_nm),
        ("lag1_definition", "Pearson correlation of residual pairs one original grid step apart; never bridge gaps"),
        ("noise_interpretation", "Exploratory fitted-residual diagnostics; cannot establish independent equal-variance measurement noise"),
        ("preprocessing_note", "Shared references, MRE concentration scaling, and smoothing can change noise correlation or variance"),
        ("source_smoothing", "Unknown: smoothing settings are not recorded by merge.py"),
    ):
        parameters.append([key, value])

    component_names = [f"SV{index + 1}" for index in range(len(svd.singular_values))]
    patterns = workbook.create_sheet(PATTERNS_SHEET)
    patterns.append(["mode", "wavelength_nm", "offset_MRE", *component_names])
    scores = workbook.create_sheet(SCORES_SHEET)
    scores.append(["mode", "spectrum", *component_names])
    for mode, decomposition in (("uncentered", analysis.uncentered), ("centered", analysis.centered)):
        for index, wavelength in enumerate(analysis.wavelengths):
            patterns.append([
                mode, float(wavelength), float(decomposition.offset[index]),
                *decomposition.patterns[:, index].tolist(),
            ])
        for name, coefficients in zip(analysis.names, decomposition.scores):
            scores.append([mode, name, *coefficients.tolist()])

    reconstructions = workbook.create_sheet(RECONSTRUCTION_SHEET)
    reconstructions.append([
        "rank", "spectrum", "wavelength_nm", "observed_MRE", "reconstructed_MRE", "residual_MRE",
    ])
    comparisons = workbook.create_sheet(COMPARISON_SHEET)
    comparisons.append([
        "rank", "spectral_energy_fraction", "relative_Frobenius_error", "residual_rms_MRE",
        "spectrum_residual_rms_CV", "wavelength_residual_rms_CV",
    ])
    noise = workbook.create_sheet(NOISE_SHEET)
    noise.append([
        "rank", "spectrum", "residual_mean_MRE", "residual_rms_MRE",
        "lag1_correlation", "lag1_pairs", "lag1_status",
    ])
    wavelength_noise = workbook.create_sheet(NOISE_WAVELENGTH_SHEET)
    wavelength_noise.append(["rank", "wavelength_nm", "residual_mean_MRE", "residual_rms_MRE"])
    for model in analysis.reconstructions:
        rms = float(np.sqrt(np.mean(model.residuals**2)))
        precision = np.finfo(float).eps * max(analysis.spectra.shape) * np.linalg.norm(analysis.spectra)
        measurable_residual = np.linalg.norm(model.residuals) > precision
        cvs = [
            float(np.std(profile) / np.mean(profile)) if measurable_residual and np.mean(profile) > 0 else None
            for profile in (model.rms_by_spectrum, model.rms_by_wavelength)
        ]
        comparisons.append([model.rank, model.energy_fraction, model.relative_error, rms, *cvs])
        comparisons.cell(comparisons.max_row, 2).number_format = "0.00%"
        for index, name in enumerate(analysis.names):
            for column, wavelength in enumerate(analysis.wavelengths):
                reconstructions.append([
                    model.rank, name, float(wavelength), float(analysis.spectra[index, column]),
                    float(model.reconstructed[index, column]), float(model.residuals[index, column]),
                ])
            correlation = float(model.lag1_correlations[index])
            if analysis.wavelength_step_nm is None:
                status = "Unavailable: original wavelength grid is irregular"
            elif model.lag1_pairs < 3:
                status = "Unavailable: fewer than three same-step residual pairs"
            elif not isfinite(correlation):
                status = "Unavailable: negligible or constant residual variation"
            else:
                status = "Estimated from fitted residuals"
            noise.append([
                model.rank, name, float(model.residuals[index].mean()), float(model.rms_by_spectrum[index]),
                correlation if isfinite(correlation) else None, model.lag1_pairs, status,
            ])
        for index, wavelength in enumerate(analysis.wavelengths):
            wavelength_noise.append([
                model.rank, float(wavelength), float(model.residuals[:, index].mean()),
                float(model.rms_by_wavelength[index]),
            ])

    for sheet in (patterns, scores, reconstructions, comparisons, noise, wavelength_noise):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float) and cell.number_format == "General":
                    cell.number_format = "0.000000E+00"


def plot_svd(analysis: MreSvdAnalysis) -> None:
    """Display patterns, scores, reconstructions, centering and residual checks."""
    import matplotlib.pyplot as plt
    from matplotlib.ticker import MaxNLocator

    svd = analysis.svd
    components = np.arange(1, len(svd.singular_values) + 1)
    figures = []

    def new_figure(*, rows=1, title, height=5):
        figure, axes = plt.subplots(rows, 2, figsize=(13, height), squeeze=False)
        figures.append(figure)
        figure.suptitle(title)
        return axes.ravel()

    def sample_axis(axis):
        positions = np.arange(1, len(analysis.names) + 1)
        axis.set_xticks(positions, labels=analysis.names, rotation=45, ha="right", fontsize=8)
        axis.set_xlabel("Spectrum (workbook order)")

    def wavelength_axis(axis, ylabel):
        axis.set_xlabel("Wavelength (nm)")
        axis.set_ylabel(ylabel)
        axis.grid(alpha=0.25)

    try:
        values_axis, energy_axis = new_figure(title=(
            f"{len(analysis.names)} MRE spectra, "
            f"{analysis.wavelengths.min():g}-{analysis.wavelengths.max():g} nm, "
            f"{len(analysis.wavelengths)} complete wavelength rows"
        ))
        values_axis.semilogy(components, svd.singular_values, "o-")
        if svd.gavish_donoho_threshold > 0:
            values_axis.axhline(
                svd.gavish_donoho_threshold, linestyle="--", label="Noise threshold",
            )
        values_axis.axvline(
            svd.gavish_donoho_rank, linestyle=":", label=f"Noise rank: {svd.gavish_donoho_rank}",
        )
        values_axis.set_title("MRE singular values")
        values_axis.set_ylabel("Singular value (MRE units, log scale)")
        values_axis.legend()

        energy_axis.bar(components, 100 * svd.explained_energy_ratio, label="Individual")
        energy_axis.plot(components, 100 * svd.cumulative_energy_ratio, "o-", label="Cumulative")
        energy_axis.axhline(
            100 * CUMULATIVE_VARIANCE_TARGET, linestyle="--",
            label=f"{100 * CUMULATIVE_VARIANCE_TARGET:g}% energy target",
        )
        energy_axis.axvline(svd.variance_rank, linestyle=":", label=f"Energy rank: {svd.variance_rank}")
        energy_axis.set_title("Explained spectral energy")
        energy_axis.set_ylabel("Spectral energy (%)")
        energy_axis.set_ylim(0, 102)
        energy_axis.legend()
        for axis in (values_axis, energy_axis):
            axis.set_xlabel("Component number")
            axis.xaxis.set_major_locator(MaxNLocator(integer=True))
            axis.grid(alpha=0.25)

        if not analysis.diagnostics_enabled:
            for figure in figures:
                figure.tight_layout()
            plt.show(block=True)
            return

        assert analysis.uncentered is not None and analysis.centered is not None
        pattern_axes = new_figure(title="Signed wavelength patterns (right singular vectors)")
        score_axes = new_figure(
            rows=analysis.diagnostic_components,
            title="Signed component scores (each component has its own scale)",
            height=max(5, 2 * analysis.diagnostic_components + 1),
        ).reshape(analysis.diagnostic_components, 2)
        positions = np.arange(1, len(analysis.names) + 1)
        for index, (mode, decomposition) in enumerate((
            ("Uncentered", analysis.uncentered), ("Mean-centered", analysis.centered),
        )):
            count = min(analysis.diagnostic_components, len(decomposition.stats.singular_values))
            has_variation = np.any(decomposition.stats.singular_values > 0)
            for component in range(count if has_variation else 0):
                label = f"SV{component + 1}"
                pattern_axes[index].plot(analysis.wavelengths, decomposition.patterns[component], label=label)
            pattern_axes[index].set_title(mode)
            pattern_axes[index].axhline(0, color="gray", linewidth=0.8)
            if has_variation:
                pattern_axes[index].legend()
            else:
                pattern_axes[index].text(
                    0.5, 0.5, "No variation after centering", transform=pattern_axes[index].transAxes, ha="center",
                )
            wavelength_axis(pattern_axes[index], "Unit-L2 loading (dimensionless)")
            for component in range(analysis.diagnostic_components):
                axis = score_axes[component, index]
                axis.set_title(f"{mode}: SV{component + 1}")
                if component < count and has_variation:
                    axis.plot(positions, decomposition.scores[:, component], "o-")
                else:
                    message = "No centered variation" if not has_variation else "No independent component"
                    axis.text(0.5, 0.5, message, transform=axis.transAxes, ha="center")
                axis.axhline(0, color="gray", linewidth=0.8)
                axis.set_ylabel("Score (MRE)")
                axis.grid(alpha=0.25)
                if component == analysis.diagnostic_components - 1:
                    sample_axis(axis)
                else:
                    axis.set_xticks(positions)
                    axis.tick_params(axis="x", labelbottom=False)

        for model in analysis.reconstructions:
            fitted_axis, residual_axis = new_figure(title=(
                f"Uncentered rank {model.rank}: {100 * model.energy_fraction:.4f}% energy, "
                f"relative residual {100 * model.relative_error:.3f}%"
            ))
            for index, name in enumerate(analysis.names):
                measured, = fitted_axis.plot(
                    analysis.wavelengths, analysis.spectra[index], linewidth=1, alpha=0.5,
                    label="Measured (solid)" if index == 0 else None,
                )
                fitted_axis.plot(
                    analysis.wavelengths, model.reconstructed[index], "--", color=measured.get_color(),
                    linewidth=1.3, label="Reconstructed (dashed)" if index == 0 else None,
                )
                residual_axis.plot(
                    analysis.wavelengths, model.residuals[index], color=measured.get_color(), label=name,
                )
            fitted_axis.set_title("Measured and reconstructed spectra")
            residual_axis.set_title("Residual = measured - reconstructed")
            fitted_axis.legend()
            residual_axis.axhline(0, color="gray", linewidth=0.8)
            if len(analysis.names) <= 6:
                residual_axis.legend(fontsize=8)
            wavelength_axis(fitted_axis, "MRE")
            wavelength_axis(residual_axis, "Residual (MRE)")

        centered_values, centered_energy = new_figure(title="Mean-centering comparison")
        for mode, decomposition in (("Uncentered", analysis.uncentered), ("Mean-centered", analysis.centered)):
            stats = decomposition.stats
            axis_components = np.arange(1, len(stats.singular_values) + 1)
            positive = stats.singular_values > 0
            centered_values.semilogy(axis_components[positive], stats.singular_values[positive], "o-", label=mode)
            if positive.any():
                centered_energy.plot(axis_components, 100 * stats.cumulative_energy_ratio, "o-", label=mode)
        centered_values.set_ylabel("Singular value (MRE units, log scale)")
        centered_values.set_title("Singular values before and after mean subtraction")
        centered_energy.set_title("Different denominators: total energy / centered variance")
        centered_energy.set_ylabel("Cumulative energy / variance (%)")
        centered_energy.axhline(100 * CUMULATIVE_VARIANCE_TARGET, linestyle="--", color="gray")
        centered_energy.set_ylim(0, 102)
        for axis in (centered_values, centered_energy):
            axis.set_xlabel("Component number")
            axis.xaxis.set_major_locator(MaxNLocator(integer=True))
            axis.grid(alpha=0.25)
            axis.legend()

        noise_axes = new_figure(rows=2, title="Fitted-residual noise diagnostics (exploratory)", height=10)
        for model in analysis.reconstructions:
            label = f"Rank {model.rank}"
            noise_axes[0].plot(positions, model.rms_by_spectrum, "o-", label=label)
            noise_axes[1].plot(analysis.wavelengths, model.rms_by_wavelength, label=label)
            if np.isfinite(model.lag1_correlations).any():
                noise_axes[2].plot(positions, model.lag1_correlations, "o-", label=label)
            noise_axes[3].plot(analysis.wavelengths, model.residuals.mean(axis=0), label=label)
        noise_axes[0].set_title("Residual size across spectra")
        noise_axes[0].set_ylabel("Residual RMS (MRE)")
        sample_axis(noise_axes[0])
        wavelength_axis(noise_axes[1], "Residual RMS (MRE)")
        noise_axes[1].set_title("Residual size across wavelength")
        if all(np.all(model.rms_by_spectrum > 0) for model in analysis.reconstructions):
            noise_axes[0].set_yscale("log")
            noise_axes[0].set_ylabel("Residual RMS (MRE, log scale)")
        if all(np.all(model.rms_by_wavelength > 0) for model in analysis.reconstructions):
            noise_axes[1].set_yscale("log")
            noise_axes[1].set_ylabel("Residual RMS (MRE, log scale)")
        noise_axes[2].set_title("One-grid-step residual correlation; gaps excluded")
        noise_axes[2].set_ylabel("Pearson correlation")
        noise_axes[2].set_ylim(-1.05, 1.05)
        sample_axis(noise_axes[2])
        if not noise_axes[2].lines:
            noise_axes[2].text(0.5, 0.5, "Correlation unavailable; see workbook status", transform=noise_axes[2].transAxes, ha="center")
        wavelength_axis(noise_axes[3], "Mean residual (MRE)")
        noise_axes[3].set_title("Mean residual across spectra")
        for axis in noise_axes:
            axis.grid(alpha=0.25)
            if axis.lines:
                axis.legend()
        for figure in figures:
            figure.tight_layout()
        plt.show(block=True)
    finally:
        for figure in figures:
            plt.close(figure)


def analyze_mre_components(
    folder: Path, *, min_nm: float = DEFAULT_MIN_NM, max_nm: float = DEFAULT_MAX_NM,
    plot: bool = False, reconstruction_ranks: list[int] | None = None, diagnostic_components: int = 5,
) -> Path:
    """Analyze the sibling combined workbook and atomically save SVD sheets."""
    folder = Path(folder).resolve()
    workbook_file = folder.parent / f"{folder.name}-combined.xlsx"
    merge_command = f'python merge.py --folder "{folder}"'
    if not workbook_file.is_file():
        raise FileNotFoundError(
            f"Combined workbook not found: {workbook_file}\nRun merge.py first: {merge_command}"
        )

    workbook = load_workbook(workbook_file, data_only=False)
    temporary_path = None
    try:
        if MRE_SHEET not in workbook.sheetnames:
            raise ValueError(
                f"{workbook_file} has no {MRE_SHEET} sheet. Run merge.py first: {merge_command}"
            )
        analysis = analyze_mre(
            workbook[MRE_SHEET], min_nm=min_nm, max_nm=max_nm,
            reconstruction_ranks=reconstruction_ranks, diagnostic_components=diagnostic_components,
        )
        write_results(workbook, analysis)
        # Preserve the source if analysis, saving, or replacement fails.
        with tempfile.NamedTemporaryFile(
            prefix=f".{workbook_file.stem}-", suffix=".xlsx", dir=workbook_file.parent, delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        workbook.close()
        os.replace(temporary_path, workbook_file)
    finally:
        workbook.close()
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)

    output_sheets = OUTPUT_SHEETS if analysis.diagnostics_enabled else BASE_SHEETS
    print(f"Saved {RESULTS_SHEET} and {len(output_sheets) - 1} supporting sheets in {workbook_file}.")
    mode = "uncentered SVD plus diagnostics" if analysis.diagnostics_enabled else "uncentered SVD only"
    print(f"Spectra: {len(analysis.names)} (saved MRE; {mode}).")
    print(
        f"Wavelength rows: {len(analysis.wavelengths)} retained, "
        f"{analysis.excluded_points} excluded for missing values, "
        f"{analysis.points_outside_range} outside the requested range."
    )
    print(f"Retained range: {analysis.wavelengths.min():g}-{analysis.wavelengths.max():g} nm.")
    print(f"Gavish-Donoho rank: {analysis.svd.gavish_donoho_rank}.")
    print(f"Rank at {100 * CUMULATIVE_VARIANCE_TARGET:g}% spectral energy: {analysis.svd.variance_rank}.")
    if analysis.diagnostics_enabled:
        assert analysis.centered is not None
        print(
            f"Mean-centered ranks: noise threshold {analysis.centered.stats.gavish_donoho_rank}, "
            f"{100 * CUMULATIVE_VARIANCE_TARGET:g}% variance {analysis.centered.stats.variance_rank} "
            f"({analysis.centered.effective_shape[0]} effective sample directions)."
        )
        for model in analysis.reconstructions:
            print(
                f"Rank-{model.rank} reconstruction: {100 * model.energy_fraction:.4f}% energy, "
                f"{100 * model.relative_error:.3f}% relative residual, "
                f"RMS {np.sqrt(np.mean(model.residuals**2)):.6g} MRE."
            )
        print("Residual noise checks are exploratory; use replicate scans or blanks to assess measurement noise.")
    print("Ranks estimate independent spectral patterns, not the number of chemical species.")
    if plot:
        plot_svd(analysis)
    return workbook_file


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Estimate spectral rank with SVD of the MRE sheet generated by merge.py."
    )
    parser.add_argument("--folder", type=Path, required=True, help="Data folder passed to merge.py.")
    parser.add_argument(
        "--min-wavelength", dest="min_nm", type=float, default=DEFAULT_MIN_NM,
        help="Lower wavelength bound, inclusive (default: 190 nm).",
    )
    parser.add_argument(
        "--max-wavelength", dest="max_nm", type=float, default=DEFAULT_MAX_NM,
        help="Upper wavelength bound, inclusive (default: 250 nm).",
    )
    parser.add_argument(
        "--plot", action="store_true",
        help="Display SVD rank plots, plus optional diagnostics when --reconstruction-ranks is supplied.",
    )
    parser.add_argument(
        "--diagnostic-components", type=int, default=5, metavar="N",
        help="Number of optional patterns/scores to plot with --reconstruction-ranks (default: 5; all are saved).",
    )
    parser.add_argument(
        "--reconstruction-ranks", nargs="+", type=int, metavar="K",
        help="Enable all additional diagnostics and compare these uncentered ranks; omit for SVD rank analysis only.",
    )
    args = parser.parse_args()
    try:
        analyze_mre_components(
            args.folder, min_nm=args.min_nm, max_nm=args.max_nm, plot=args.plot,
            reconstruction_ranks=args.reconstruction_ranks, diagnostic_components=args.diagnostic_components,
        )
    except (OSError, ValueError, RuntimeError, BadZipFile, np.linalg.LinAlgError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
