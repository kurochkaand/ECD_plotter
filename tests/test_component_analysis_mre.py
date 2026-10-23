"""Check MRE SVD, common-wavelength filtering, and safe workbook updates."""

import io
import os
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

import matplotlib
import numpy as np
from openpyxl import Workbook, load_workbook

matplotlib.use("Agg")
import matplotlib.pyplot as plt

import component_analysis_mre as analysis
import spectral_component_analysis as original_analysis


PROJECT_ROOT = Path(__file__).resolve().parents[1]


class MreSvdTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="ECD MRE SVD ")
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name).resolve() / "raw data"
        self.folder.mkdir()
        self.workbook_file = self.folder.parent / "raw data-combined.xlsx"
        self.wavelengths = np.arange(185, 256, dtype=float)
        components = np.vstack([
            -10000 * np.exp(-((self.wavelengths - 208) / 10) ** 2),
            8000 * np.exp(-((self.wavelengths - 235) / 9) ** 2),
        ])
        fraction = np.linspace(0.05, 0.95, 8)
        self.spectra = np.column_stack([fraction, 1 - fraction]) @ components
        # A noise estimator needs a noise floor above floating-point roundoff.
        self.spectra += np.random.default_rng(17).normal(0, 0.1, self.spectra.shape)
        self.names = [f"sample {index + 1}" for index in range(len(fraction))]
        self.write_workbook()

    def write_workbook(self, *, wavelengths=None, spectra=None, names=None):
        wavelengths = self.wavelengths if wavelengths is None else wavelengths
        spectra = self.spectra if spectra is None else spectra
        names = self.names if names is None else names
        workbook = Workbook()
        raw = workbook.active
        raw.title = "CD_subtr"
        raw.append(["measured", "reference", "difference"])
        raw.append([5, 2, "=A2-B2"])
        source = workbook.create_sheet("MRE")
        source.append(["wavelength_nm", *names])
        for wavelength, signals in zip(wavelengths, spectra.T):
            source.append([float(wavelength), *signals.tolist()])
        for title in ("MRE_parameters", "SESCA_results", "SESCA_fit", "Epsilon"):
            workbook.create_sheet(title).append(["keep", title])
        workbook.save(self.workbook_file)
        workbook.close()

    def open_workbook(self):
        workbook = load_workbook(self.workbook_file, data_only=False)
        self.addCleanup(workbook.close)
        return workbook

    def run_analysis(self, **kwargs):
        with redirect_stdout(io.StringIO()) as output:
            result = analysis.analyze_mre_components(self.folder, **kwargs)
        return result, output.getvalue()

    def parameters(self, workbook):
        return dict(list(workbook[analysis.PARAMETERS_SHEET].values)[1:])

    def run_cli(self, *arguments):
        return subprocess.run(
            [sys.executable, str(PROJECT_ROOT / "component_analysis_mre.py"),
             "--folder", str(self.folder) + "\\", *arguments],
            env={**os.environ, "MPLBACKEND": "Agg"},
            capture_output=True, text=True, check=False,
        )

    def test_known_two_component_signed_mre_uses_defaults_without_raw_files_or_table(self):
        self.assertEqual(list(self.folder.iterdir()), [])
        with patch.object(analysis, "plot_svd") as plot:
            result, output = self.run_analysis()
        plot.assert_not_called()
        self.assertEqual(result, self.workbook_file)
        workbook = self.open_workbook()
        parameters = self.parameters(workbook)
        self.assertEqual(parameters["gavish_donoho_rank"], 2)
        self.assertEqual(parameters["spectral_energy_rank"], 2)
        self.assertEqual(parameters["requested_min_nm"], 190)
        self.assertEqual(parameters["requested_max_nm"], 250)
        self.assertEqual(parameters["retained_wavelength_rows"], 61)
        self.assertEqual(parameters["excluded_missing_rows"], 0)
        self.assertEqual(parameters["outside_range_rows"], 10)
        self.assertFalse(parameters["mean_centered"])
        self.assertIn("Gavish-Donoho rank: 2", output)
        fractions = [row[2] for row in list(workbook[analysis.RESULTS_SHEET].values)[1:]]
        self.assertAlmostEqual(sum(fractions), 1)

    def test_saved_input_preserves_mre_units_and_original_sheets_and_formulas(self):
        source = self.open_workbook()
        original = {sheet.title: list(sheet.values) for sheet in source}
        source.close()
        self.run_analysis()
        workbook = self.open_workbook()
        for title, values in original.items():
            self.assertEqual(list(workbook[title].values), values)
        inputs = list(workbook[analysis.INPUT_SHEET].values)
        self.assertEqual(inputs[0], ("wavelength_nm", *self.names))
        expected = (self.wavelengths >= 190) & (self.wavelengths <= 250)
        np.testing.assert_allclose(np.asarray(inputs[1:])[:, 1:], self.spectra[:, expected].T)
        self.assertEqual(workbook["CD_subtr"]["C2"].value, "=A2-B2")

    def test_blanks_exclude_entire_rows_and_descending_input_is_sorted(self):
        spectra = self.spectra.astype(object)
        for wavelength, spectrum in ((190, 0), (201, 1), (250, 2), (185, 3)):
            spectra[spectrum, np.flatnonzero(self.wavelengths == wavelength)[0]] = None
        self.write_workbook(wavelengths=self.wavelengths[::-1], spectra=spectra[:, ::-1])
        self.run_analysis()
        workbook = self.open_workbook()
        parameters = self.parameters(workbook)
        self.assertEqual(parameters["excluded_missing_rows"], 3)
        self.assertEqual(parameters["retained_wavelength_rows"], 58)
        self.assertEqual(parameters["retained_min_nm"], 191)
        self.assertEqual(parameters["retained_max_nm"], 249)
        inputs = np.asarray(list(workbook[analysis.INPUT_SHEET].values)[1:], dtype=float)
        self.assertTrue(np.isfinite(inputs).all())
        self.assertTrue(np.all(np.diff(inputs[:, 0]) > 0))
        self.assertNotIn(201, inputs[:, 0])

    def test_rank_is_uncentered_even_if_old_script_centering_is_enabled(self):
        workbook = self.open_workbook()
        with patch.object(original_analysis, "MEAN_CENTER_FOR_RANK", True):
            result = analysis.analyze_mre(workbook["MRE"])
            centered = original_analysis.analyze_singular_values(result.spectra)
        self.assertEqual(result.svd.variance_rank, 2)
        self.assertEqual(centered.variance_rank, 1)

    def test_all_singular_values_and_energy_rank_are_reported_beyond_four(self):
        wavelengths = np.arange(190, 251, dtype=float)
        spectra = np.vstack([
            10000 * np.sin(np.pi * index * (wavelengths - 190) / 60)
            for index in range(1, 7)
        ])
        self.write_workbook(wavelengths=wavelengths, spectra=spectra, names=self.names[:6])
        self.run_analysis()
        workbook = self.open_workbook()
        self.assertEqual(workbook[analysis.RESULTS_SHEET].max_row, 7)
        self.assertEqual(self.parameters(workbook)["spectral_energy_rank"], 6)

    def test_rerun_replaces_only_svd_outputs_and_updates_range(self):
        self.run_analysis()
        self.run_analysis(min_nm=200, max_nm=220)
        workbook = self.open_workbook()
        for title in (analysis.RESULTS_SHEET, analysis.PARAMETERS_SHEET, analysis.INPUT_SHEET):
            self.assertEqual(workbook.sheetnames.count(title), 1)
        parameters = self.parameters(workbook)
        self.assertEqual(parameters["retained_wavelength_rows"], 21)
        self.assertEqual(parameters["requested_min_nm"], 200)
        self.assertEqual(parameters["requested_max_nm"], 220)
        self.assertIn("SESCA_results", workbook.sheetnames)
        self.assertEqual(workbook["MRE"].max_row, 72)

    def test_plot_uses_full_diagnostics_and_runs_after_successful_save(self):
        def inspect_plot(*args, **kwargs):
            saved = self.open_workbook()
            self.assertIn(analysis.RESULTS_SHEET, saved.sheetnames)
            self.assertEqual(len(plt.gcf().axes), 2)
            np.testing.assert_array_equal(plt.gcf().axes[0].lines[0].get_xdata(), np.arange(1, 9))
            self.assertIn("190-250 nm", plt.gcf()._suptitle.get_text())

        with patch.object(plt, "show", side_effect=inspect_plot) as show:
            self.run_analysis(plot=True)
        show.assert_called_once_with(block=True)
        self.assertEqual(plt.get_fignums(), [])

    def test_cli_supports_folder_with_spaces_trailing_separator_and_range_overrides(self):
        result = self.run_cli("--min-wavelength", "200", "--max-wavelength", "220")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("200-220 nm", result.stdout)
        self.assertIn("MRE_SVD_results", result.stdout)
        self.assertEqual(self.parameters(self.open_workbook())["retained_wavelength_rows"], 21)

    def test_invalid_ranges_preserve_existing_outputs(self):
        self.run_analysis()
        original = self.workbook_file.read_bytes()
        for minimum, maximum in ((250, 190), (190, 190), (0, 250), (np.nan, 250), (190, np.inf)):
            with self.subTest(minimum=minimum, maximum=maximum):
                with self.assertRaisesRegex(ValueError, "bounds"):
                    self.run_analysis(min_nm=minimum, max_nm=maximum)
                self.assertEqual(self.workbook_file.read_bytes(), original)

    def test_outside_range_and_insufficient_complete_rows_preserve_workbook(self):
        original = self.workbook_file.read_bytes()
        with self.assertRaisesRegex(ValueError, "workbook spans 185-255"):
            self.run_analysis(min_nm=300, max_nm=400)
        self.assertEqual(self.workbook_file.read_bytes(), original)

        spectra = self.spectra.astype(object)
        spectra[0, :] = None
        self.write_workbook(spectra=spectra)
        original = self.workbook_file.read_bytes()
        with self.assertRaisesRegex(ValueError, "Only 0 complete.*excluding 61"):
            self.run_analysis()
        self.assertEqual(self.workbook_file.read_bytes(), original)

    def test_zero_data_and_single_spectrum_are_rejected(self):
        for spectra, names, message in (
            (np.zeros_like(self.spectra), self.names, "zero numerical rank"),
            (self.spectra[:1], self.names[:1], "at least two spectra"),
        ):
            with self.subTest(message=message):
                self.write_workbook(spectra=spectra, names=names)
                original = self.workbook_file.read_bytes()
                with self.assertRaisesRegex(ValueError, message):
                    self.run_analysis()
                self.assertEqual(self.workbook_file.read_bytes(), original)

    def test_invalid_mre_values_layout_and_duplicate_wavelengths_are_rejected(self):
        for cell, value, message in (
            ("B7", "=B6", "MRE!B7"),
            ("B7", "#DIV/0!", "MRE!B7"),
            ("B7", "bad MRE", "MRE!B7"),
            ("A1", "wrong", "wavelength_nm"),
            ("C1", self.names[0], "Duplicate spectrum names"),
            ("B1", " ", "nonempty names"),
            ("A8", 190, "wavelengths must be unique"),
        ):
            with self.subTest(cell=cell, value=value):
                self.write_workbook()
                workbook = self.open_workbook()
                workbook["MRE"][cell] = value
                workbook.save(self.workbook_file)
                workbook.close()
                original = self.workbook_file.read_bytes()
                with self.assertRaisesRegex(ValueError, message):
                    self.run_analysis()
                self.assertEqual(self.workbook_file.read_bytes(), original)

    def test_missing_mre_or_workbook_reports_merge_command_without_traceback(self):
        workbook = self.open_workbook()
        del workbook["MRE"]
        workbook.save(self.workbook_file)
        workbook.close()
        original = self.workbook_file.read_bytes()
        result = self.run_cli()
        self.assertEqual(result.returncode, 1)
        self.assertIn("no MRE sheet", result.stderr)
        self.assertIn("Run merge.py first", result.stderr)
        self.assertNotIn("Traceback", result.stderr)
        self.assertEqual(self.workbook_file.read_bytes(), original)
        self.workbook_file.unlink()
        result = self.run_cli()
        self.assertEqual(result.returncode, 1)
        self.assertIn(str(self.workbook_file), result.stderr)
        self.assertIn("Run merge.py first", result.stderr)
        self.assertNotIn("Traceback", result.stderr)

    def test_failed_save_or_replacement_preserves_source_and_cleans_temporary_files(self):
        for target in ("save", "replace"):
            with self.subTest(target=target):
                original = self.workbook_file.read_bytes()
                owner = analysis.os if target == "replace" else Workbook
                with patch.object(owner, target, side_effect=PermissionError("Workbook is open")):
                    with self.assertRaises(PermissionError):
                        self.run_analysis()
                self.assertEqual(self.workbook_file.read_bytes(), original)
                self.assertEqual(list(self.workbook_file.parent.glob(".*-combined-*.xlsx")), [])


if __name__ == "__main__":
    unittest.main()
