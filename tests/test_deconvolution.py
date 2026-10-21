"""Check CD unit conversion, merge integration, and safe workbook updates."""

import io
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import deconvolution
import merge


PROJECT_ROOT = Path(__file__).resolve().parents[1]


class MreConversionTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="ECD MRE ")
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name).resolve() / "sample data"
        self.folder.mkdir()
        self.write_spectrum("peptide", [-20, 0, 20])
        self.write_table("peptide\t0.1\n")

    def write_table(self, rows, *, pathlength=1, residues=20, header=None):
        header = header or "spectrum\tconcentration_in_mM"
        (self.folder / "concentration_table.txt").write_text(
            f"pathlength_in_mm = {pathlength}\nnumber_of_amino_acids = {residues}\n\n"
            + header + "\n" + rows, encoding="utf-8",
        )

    def write_spectrum(self, name, values, ht=(400, 400, 400)):
        rows = "".join(
            f"{wavelength},{signal},{voltage},0.1\n"
            for wavelength, signal, voltage in zip((200, 201, 202), values, ht)
        )
        (self.folder / f"{name}.csv").write_text(
            "TITLE,Instrument export\nDATA TYPE,\nORIGIN,JASCO\nOWNER,\n"
            "DATE,26/10/14\nTIME,18:41:36\nXYDATA\n" + rows
            + "##### Extended Information\n", encoding="latin-1",
        )

    def open_workbook(self, path, **kwargs):
        workbook = load_workbook(path, **kwargs)
        self.addCleanup(workbook.close)
        return workbook

    def test_conversion_preserves_raw_and_normalized_mdeg(self):
        self.write_table("peptide\t0.2\n")
        output = merge.merge_spectra(self.folder)
        workbook = self.open_workbook(output)
        self.assertEqual([row[1] for row in list(workbook["CD"].values)[1:]], [-20, 0, 20])
        self.assertEqual([row[3] for row in list(workbook["CD_subtr"].values)[1:]], [-100, 0, 100])
        # -20 mdeg / (1 mm * 0.0002 M * 20 residues) = -5000.
        self.assertEqual(list(workbook["MRE"].values), [
            ("wavelength_nm", "peptide"), (200, -5000), (201, 0), (202, 5000),
        ])
        metadata = dict(list(workbook["MRE_parameters"].values)[1:])
        self.assertEqual(metadata["concentration_in_mM[peptide]"], 0.2)
        self.assertEqual(metadata["concentration_in_M[peptide]"], 0.0002)
        self.assertEqual(metadata["pathlength_in_mm"], 1)
        self.assertEqual(metadata["number_of_amino_acids"], 20)
        self.assertIn("dmol^-1", metadata["MRE_units"])

    def test_subtraction_works_without_excel_cached_results(self):
        self.write_spectrum("H2O", [2, 4, 6])
        self.write_spectrum("second", [10, 12, 14])
        self.write_table("H2O\t1\npeptide\t0.05\nsecond\t0.1\n")
        output = merge.merge_spectra(self.folder, subtract="H2O")
        cached = self.open_workbook(output, data_only=True)
        self.assertIsNone(cached["CD_subtr"]["F2"].value)
        workbook = self.open_workbook(output)
        self.assertEqual(workbook["CD_subtr"]["F2"].value, "=(C2-$B2)/0.05")
        self.assertEqual(workbook["CD_subtr"]["G2"].value, "=(D2-$B2)/0.1")
        self.assertEqual(list(workbook["MRE"].values), [
            ("wavelength_nm", "peptide", "second"),
            (200, -22000, 4000), (201, -4000, 4000), (202, 14000, 4000),
        ])

    def test_ht_exclusions_remain_blank_with_subtraction_and_smoothing(self):
        self.write_spectrum("peptide", [-20, 0, 20], ht=(600, 400, 400))
        self.write_spectrum("H2O", [2, 4, 6], ht=(400, 600, 400))
        output = merge.merge_spectra(self.folder, subtract="H2O", ht_threshold=500, smooth=3)
        workbook = self.open_workbook(output)
        self.assertEqual(list(workbook["MRE"].values), [
            ("wavelength_nm", "peptide"), (200, None), (201, None), (202, 7000),
        ])

    def test_table_edits_reach_mre_after_rerunning_merge(self):
        self.write_spectrum("H2O", [2, 4, 6])
        self.write_spectrum("removed", [10, 12, 14])
        self.write_table("peptide\t0.1\nremoved\t0.1\n")
        output = merge.merge_spectra(self.folder, subtract="H2O")

        table_file = self.folder / "concentration_table.txt"
        self.write_table(
            "peptide\t\t0.2\t7\tKept sample\n",
            header="spectrum\t\tconcentration_in_mM\tPh\tshort name",
        )
        original = table_file.read_bytes()
        merge.merge_spectra(self.folder, subtract="H2O")

        workbook = self.open_workbook(output)
        self.assertEqual(list(workbook["MRE"].values), [
            ("wavelength_nm", "Kept sample"),
            (200, -5500), (201, -1000), (202, 3500),
        ])
        self.assertEqual(table_file.read_bytes(), original)

    def test_rerun_replaces_outputs_with_updated_parameters(self):
        output = merge.merge_spectra(self.folder)
        self.write_table("peptide\t0.1\n", pathlength=2)
        merge.merge_spectra(self.folder)
        workbook = self.open_workbook(output)
        self.assertEqual(workbook.sheetnames.count("MRE"), 1)
        self.assertEqual(workbook.sheetnames.count("MRE_parameters"), 1)
        self.assertEqual(workbook["MRE"]["B2"].value, -5000)
        self.assertEqual(dict(list(workbook["MRE_parameters"].values)[1:])["pathlength_in_mm"], 2)

    def test_residue_count_edits_recalculate_mre(self):
        output = merge.merge_spectra(self.folder)
        self.write_table("peptide\t0.1\n", residues=40)
        merge.merge_spectra(self.folder)
        workbook = self.open_workbook(output)
        self.assertEqual(workbook["MRE"]["B2"].value, -5000)
        self.assertEqual(dict(list(workbook["MRE_parameters"].values)[1:])["number_of_amino_acids"], 40)

    def test_fully_excluded_spectra_keep_blank_columns(self):
        self.write_spectrum("peptide", [-20, 0, 20], ht=(600, 600, 600))
        output = merge.merge_spectra(self.folder, ht_threshold=500)
        workbook = self.open_workbook(output)
        self.assertEqual(list(workbook["MRE"].values), [
            ("wavelength_nm", "peptide"), (200, None), (201, None), (202, None),
        ])

    def test_missing_mre_sheet_suggests_merge_with_table_settings(self):
        output = merge.merge_spectra(self.folder)
        workbook = self.open_workbook(output)
        del workbook["MRE"]
        workbook.save(output)
        workbook.close()
        original_bytes = output.read_bytes()
        result = subprocess.run(
            [sys.executable, str(PROJECT_ROOT / "deconvolution.py"), "--folder", str(self.folder)],
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(result.returncode, 1)
        self.assertIn("no MRE sheet", result.stderr)
        self.assertIn("concentration_table.txt", result.stderr)
        self.assertNotIn("--concentration-in-mM", result.stderr)
        self.assertEqual(output.read_bytes(), original_bytes)

    def test_invalid_parameters_leave_workbook_unchanged(self):
        output = merge.merge_spectra(self.folder)
        original_bytes = output.read_bytes()
        for parameter, values in (
            ("pathlength", [0, -1, "nan", "inf", "bad"]),
            ("concentration", [0, -1, "nan", "inf", "bad"]),
            ("residues", [0, -1, 2.5, True, "bad"]),
        ):
            for value in values:
                with self.subTest(parameter=parameter, value=value):
                    parameters = {parameter: value} if parameter != "concentration" else {}
                    concentration = value if parameter == "concentration" else 0.1
                    self.write_table(f"peptide\t{concentration}\n", **parameters)
                    original_table = (self.folder / "concentration_table.txt").read_bytes()
                    with self.assertRaises(ValueError):
                        merge.merge_spectra(self.folder)
                    self.assertEqual((self.folder / "concentration_table.txt").read_bytes(), original_table)
                    self.assertEqual(output.read_bytes(), original_bytes)

    def test_invalid_data_and_unsupported_formulas_leave_workbook_unchanged(self):
        for value in ("bad CD", "#DIV/0!", "=SUM(B2)", "=B2-$B3", "=(B2-$B2)/0.0"):
            with self.subTest(value=value):
                output = merge.merge_spectra(self.folder)
                workbook = self.open_workbook(output)
                workbook["CD_subtr"]["D2"] = value
                workbook.save(output)
                original_bytes = output.read_bytes()
                with self.assertRaisesRegex(ValueError, "CD_subtr!D2"):
                    merge.write_mre_spectra(workbook, path_in_mm=1, num_of_aa=20, concentration_in_M=0.001)
                self.assertEqual(output.read_bytes(), original_bytes)

    def test_failed_replace_preserves_source_and_cleans_temporary_workbook(self):
        output = merge.merge_spectra(self.folder)
        original_bytes = output.read_bytes()
        with patch.object(merge.os, "replace", side_effect=PermissionError("Workbook is open")):
            with self.assertRaises(PermissionError):
                merge.merge_spectra(self.folder)
        self.assertEqual(output.read_bytes(), original_bytes)
        self.assertEqual(list(output.parent.glob(".*-combined-*.xlsx")), [])

    def run_cli(self, folder=None):
        return subprocess.run(
            [sys.executable, str(PROJECT_ROOT / "merge.py"),
             "--folder", str(folder or self.folder)],
            capture_output=True, text=True, check=False,
        )

    def test_cli_accepts_millimolar_concentration_and_folder_with_spaces(self):
        output = merge.merge_spectra(self.folder)
        result = self.run_cli()
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Saved MRE", result.stdout)
        self.assertEqual(self.open_workbook(output, data_only=True)["MRE"]["B2"].value, -10000)

    def test_cli_missing_workbook_suggests_merge(self):
        result = subprocess.run(
            [sys.executable, str(PROJECT_ROOT / "deconvolution.py"), "--folder", str(self.folder)],
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(result.returncode, 1)
        self.assertIn(str(self.folder.parent / "sample data-combined.xlsx"), result.stderr)
        self.assertIn("Run merge.py first", result.stderr)
        self.assertNotIn("Traceback", result.stderr)

    def test_removed_mre_flags_are_rejected_before_writing(self):
        output = merge.merge_spectra(self.folder)
        original_bytes = output.read_bytes()
        for flag in ("--path-in-mm", "--num-of-aa", "--concentration-in-mM"):
            with self.subTest(flag=flag):
                with patch.object(sys, "argv", ["merge.py", "--folder", str(self.folder), flag, "1"]):
                    with patch("sys.stderr", io.StringIO()) as stderr, self.assertRaises(SystemExit) as error:
                        merge.main()
                self.assertEqual(error.exception.code, 2)
                self.assertIn("unrecognized arguments", stderr.getvalue())
                self.assertEqual(output.read_bytes(), original_bytes)

    def test_old_or_incomplete_tables_report_missing_settings_without_writing(self):
        output = merge.merge_spectra(self.folder)
        original_bytes = output.read_bytes()
        table_file = self.folder / "concentration_table.txt"
        for contents in (
            "spectrum\tconcentration_coefficient\npeptide\t1\n",
            "spectrum\tconcentration_in_mM\npeptide\t0.1\n",
            "pathlength_in_mm = 1\nspectrum\tconcentration_in_mM\npeptide\t0.1\n",
            "number_of_amino_acids = 20\nspectrum\tconcentration_in_mM\npeptide\t0.1\n",
        ):
            with self.subTest(contents=contents):
                table_file.write_text(contents, encoding="utf-8")
                original_table = table_file.read_bytes()
                with self.assertRaisesRegex(ValueError, "missing MRE settings"):
                    merge.merge_spectra(self.folder)
                self.assertEqual(table_file.read_bytes(), original_table)
                self.assertEqual(output.read_bytes(), original_bytes)
        result = self.run_cli()
        self.assertEqual(result.returncode, 1)
        self.assertIn("missing MRE settings", result.stderr)
        self.assertIn("pathlength_in_mm", result.stderr)
        self.assertNotIn("Traceback", result.stderr)

    def test_old_concentration_header_is_rejected_even_with_new_settings(self):
        self.write_table("peptide\t1\n", header="spectrum\tconcentration_coefficient")
        with self.assertRaisesRegex(ValueError, "Rename concentration_coefficient"):
            merge.merge_spectra(self.folder)

    def test_unknown_and_duplicate_settings_are_rejected(self):
        table_file = self.folder / "concentration_table.txt"
        for extra, message in (
            ("pathlength_in_mm = 2\n", "Duplicate MRE setting"),
            ("pathlength_in_cm = 1\n", "Unknown MRE setting"),
        ):
            with self.subTest(extra=extra):
                self.write_table("peptide\t0.1\n")
                contents = table_file.read_text(encoding="utf-8")
                table_file.write_text(extra + contents, encoding="utf-8")
                with self.assertRaisesRegex(ValueError, message):
                    merge.merge_spectra(self.folder)

    def test_invalid_mre_values_are_rejected_without_modifying_workbook(self):
        for value in ("bad MRE", "#DIV/0!", "=B2*1000"):
            with self.subTest(value=value):
                output = merge.merge_spectra(self.folder)
                workbook = self.open_workbook(output)
                workbook["MRE"]["B2"] = value
                workbook.save(output)
                original_bytes = output.read_bytes()
                with self.assertRaisesRegex(ValueError, "MRE!B2"):
                    deconvolution.deconvolve_spectra(self.folder)
                self.assertEqual(output.read_bytes(), original_bytes)

    def test_invalid_mre_layout_is_rejected_without_modifying_workbook(self):
        for problem in ("wavelength_header", "no_spectra", "no_rows", "blank_name", "duplicate_names"):
            with self.subTest(problem=problem):
                output = merge.merge_spectra(self.folder)
                workbook = self.open_workbook(output)
                sheet = workbook["MRE"]
                if problem == "wavelength_header":
                    sheet["A1"] = "wrong"
                elif problem == "no_spectra":
                    sheet.delete_cols(2)
                elif problem == "no_rows":
                    sheet.delete_rows(2, 3)
                elif problem == "blank_name":
                    sheet["B1"] = " "
                else:
                    sheet["C1"] = "peptide"
                workbook.save(output)
                original_bytes = output.read_bytes()
                with self.assertRaisesRegex(ValueError, "MRE"):
                    deconvolution.deconvolve_spectra(self.folder)
                self.assertEqual(output.read_bytes(), original_bytes)

    def test_conversion_overflow_preserves_existing_workbook(self):
        output = merge.merge_spectra(self.folder)
        original_bytes = output.read_bytes()
        self.write_table("peptide\t1e-306\n")
        with self.assertRaisesRegex(ValueError, "Nonfinite MRE"):
            merge.merge_spectra(self.folder)
        self.assertEqual(output.read_bytes(), original_bytes)
        self.assertEqual(list(output.parent.glob(".*-combined-*.xlsx")), [])


if __name__ == "__main__":
    unittest.main()
