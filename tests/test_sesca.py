"""Verify SESCA units, missing-data handling, and the real upstream engine."""

import io
import random
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import matplotlib
import numpy as np
from openpyxl import Workbook, load_workbook

matplotlib.use("Agg")
import matplotlib.pyplot as plt

import deconvolution
import sesca_integration as sesca


PROJECT_ROOT = Path(__file__).resolve().parents[1]


class FakeEngine:
    """Small API double that records the actual kMRE input passed to SESCA."""

    def __init__(self, basis):
        self.basis = {row[0]: row[1:] for row in basis}
        self.Def_Args = []
        self.inputs = []
        self.arguments = []

    def Pass_Defaults(self):
        return self.Def_Args

    def Read_Args(self, arguments):
        if self.Def_Args:
            raise AssertionError("Defaults were not reset between spectra")
        self.Def_Args.append("mutated")
        self.arguments.append(arguments)
        return arguments

    def Deconv_Main(self, arguments):
        values = np.loadtxt(arguments[arguments.index("@spect") + 1])
        self.inputs.append(values)
        fractions = [0.2, 0.3, 0.5]
        predicted = np.asarray([self.basis[wave] for wave in values[:, 0]]) @ fractions
        rmsd = float(np.sqrt(np.mean((values[:, 1] - predicted) ** 2)))
        amplitude = float(np.sqrt(np.mean(values[:, 1] ** 2)))
        return [fractions, rmsd, 0.0, rmsd / amplitude, 1.0, [], ""]


class SescaTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="SESCA adapter ")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name).resolve()
        (self.root / "scripts").mkdir()
        (self.root / "scripts" / "SESCA_deconv.py").write_text("# Fake for adapter tests\n")
        (self.root / "libs").mkdir()
        self.basis = np.asarray([
            [190 + 2 * i, np.sin(i * 0.2) - 2, np.cos(i * 0.35), i / 10 + 0.5]
            for i in range(16)
        ])
        basis_file = self.root / "libs" / "Map_DS-dT.dat"
        np.savetxt(basis_file, self.basis, header="| Wavelength Alpha Beta Coil", comments="#", fmt="%.15g")
        self.observed = self.basis[:, 1:] @ [0.2, 0.3, 0.5] * 1000
        self.rows = [[wave, float(value)] for wave, value in zip(self.basis[:, 0], self.observed)]
        self.options = sesca.SescaOptions(directory=self.root)
        self.engine = FakeEngine(self.basis)
        loader = patch.object(sesca, "load_engine", return_value=self.engine)
        loader.start()
        self.addCleanup(loader.stop)

    def write_combined(self):
        folder = self.root / "spectra"
        folder.mkdir()
        output = self.root / "spectra-combined.xlsx"
        workbook = Workbook()
        source = workbook.active
        source.title = "CD_subtr"
        source.append(["wavelength_nm", "sample", "zero", "sample"])
        for wave, value in self.rows:
            raw = value * 0.002  # 1 mm * 0.0001 M * 20 residues
            source.append([wave, raw, 0, raw])
        workbook.create_sheet("Notes")["A1"] = "Preserve"
        workbook.save(output)
        workbook.close()
        return folder, output

    def test_mre_is_converted_to_kmre_and_input_defaults_are_reset(self):
        rows = [[wave, value, value] for wave, value in self.rows]
        result = sesca.fit_spectra(["one", "two"], rows, self.options)
        self.assertEqual([fit.status for fit in result.fits], ["ok", "ok"])
        self.assertEqual(len(self.engine.inputs), 2)
        np.testing.assert_allclose(self.engine.inputs[0][:, 1], self.observed / 1000, atol=1e-10)
        np.testing.assert_allclose(result.fits[0].fitted_mre, self.observed, atol=1e-10)
        self.assertEqual(result.fits[0].fractions, [0.2, 0.3, 0.5])
        self.assertEqual(dict(result.parameters)["sesca_input_units"], "kMRE = MRE / 1000")

    def test_excluded_grid_points_are_never_filled_from_neighboring_measurements(self):
        self.rows[4][1] = None
        self.rows.extend([[198.1, 99999], [197.9, 99999]])
        result = sesca.fit_spectra(["sample"], self.rows, self.options)
        self.assertEqual(len(result.fits[0].wavelengths), 15)
        self.assertNotIn(198, result.fits[0].wavelengths)
        self.assertNotIn(198, self.engine.inputs[0][:, 0])

    def test_blank_zero_sparse_and_narrow_spectra_are_reported_as_skipped(self):
        rows = [[wave, None, 0.0, value if wave < 200 else None,
                 value if 198 <= wave <= 216 else None] for wave, value in self.rows]
        result = sesca.fit_spectra(["blank", "zero", "sparse", "narrow"], rows, self.options)
        self.assertEqual([fit.status for fit in result.fits], ["skipped"] * 4)
        self.assertTrue(all(not fit.fractions and fit.note for fit in result.fits))
        self.assertEqual(self.engine.inputs, [])

    def test_invalid_settings_and_duplicate_wavelengths_are_rejected(self):
        for changes in ({"min_nm": 250}, {"max_nm": float("nan")}, {"repeats": 0}, {"seed": -1}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                sesca.fit_spectra(["sample"], self.rows, sesca.SescaOptions(directory=self.root, **changes))
        with self.assertRaisesRegex(ValueError, "unique wavelengths"):
            sesca.fit_spectra(["sample"], self.rows + [self.rows[0]], self.options)

    def test_missing_installation_and_basis_have_actionable_errors(self):
        with self.assertRaisesRegex(FileNotFoundError, "--sesca-dir.*--mre-only"):
            sesca.fit_spectra(["sample"], self.rows, sesca.SescaOptions(directory=self.root / "missing"))
        with self.assertRaisesRegex(FileNotFoundError, "--sesca-basis"):
            sesca.fit_spectra(["sample"], self.rows, sesca.SescaOptions(directory=self.root, basis="missing"))

    def test_workbook_outputs_and_mre_only_reruns_remove_stale_fits(self):
        folder, output = self.write_combined()
        parameters = dict(path_in_mm=1, concentration_in_M=0.0001, num_of_aa=20)
        deconvolution.convert_to_mre(folder, **parameters, sesca_options=self.options)
        workbook = load_workbook(output)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["Notes"]["A1"].value, "Preserve")
        self.assertEqual(list(workbook["SESCA_results"].values)[1][0:6], ("sample", "ok", None, 0.2, 0.3, 0.5))
        self.assertEqual(workbook["SESCA_results"]["D2"].number_format, "0.00%")
        self.assertEqual(workbook["SESCA_fit"].max_row, 17)
        first = list(workbook["SESCA_fit"].values)[1]
        self.assertAlmostEqual(first[2], first[3])
        self.assertAlmostEqual(first[3], first[4])
        self.assertAlmostEqual(first[5], 0, places=7)
        deconvolution.convert_to_mre(folder, **parameters)
        rerun = load_workbook(output)
        self.addCleanup(rerun.close)
        self.assertNotIn("SESCA_results", rerun.sheetnames)
        self.assertNotIn("SESCA_fit", rerun.sheetnames)

    def test_engine_failure_preserves_original_workbook(self):
        folder, output = self.write_combined()
        original = output.read_bytes()
        with patch.object(self.engine, "Deconv_Main", side_effect=SystemExit(2)):
            with self.assertRaisesRegex(RuntimeError, "SESCA failed for 'sample'"):
                deconvolution.convert_to_mre(
                    folder, path_in_mm=1, concentration_in_M=0.0001, num_of_aa=20,
                    sesca_options=self.options,
                )
        self.assertEqual(output.read_bytes(), original)

    def test_random_generator_states_are_restored(self):
        python_state, numpy_state = random.getstate(), np.random.get_state()
        sesca.fit_spectra(["sample"], self.rows, self.options)
        self.assertEqual(random.getstate(), python_state)
        new_state = np.random.get_state()
        self.assertEqual(numpy_state[0], new_state[0])
        np.testing.assert_array_equal(numpy_state[1], new_state[1])
        self.assertEqual(numpy_state[2:], new_state[2:])

    def test_plot_is_disabled_by_default(self):
        folder, _ = self.write_combined()
        with patch.object(deconvolution, "plot_sesca_results") as plot:
            deconvolution.convert_to_mre(
                folder, path_in_mm=1, concentration_in_M=0.0001, num_of_aa=20,
                sesca_options=self.options,
            )
        plot.assert_not_called()

    def test_plot_reads_saved_class_fractions_in_spectrum_order(self):
        output = self.root / "results.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "SESCA_results"
        classes = ["Helix1", "Beta1", "Helix2", "Turn1", "Turn2", "Other"]
        sheet.append(["spectrum", "status", *[f"{name}_fraction" for name in classes], "RMSD_kMRE"])
        fractions = np.asarray([[0.1, 0.2, 0.15, 0.25, 0.1, 0.2], [0] * 6, [0.5, 0.1, 0, 0, 0.1, 0.3]])
        sheet.append(["first sample", "ok", *fractions[0], 0.4])
        sheet.append(["excluded sample", "skipped", *[None] * 6, None])
        sheet.append(["last sample", "ok", *fractions[2], 0.8])
        workbook.save(output)
        workbook.close()
        figures = []

        def inspect_chart(*args, **kwargs):
            figure = plt.gcf()
            figures.append(figure)
            axis = figure.axes[0]
            self.assertEqual(axis.get_xlabel(), "Secondary structure (%)")
            self.assertEqual(axis.get_xlim(), (0, 100))
            self.assertEqual([tick.get_text() for tick in axis.get_yticklabels()], [
                "first sample", "excluded sample", "last sample",
            ])
            self.assertTrue(axis.yaxis_inverted())
            self.assertEqual([item.get_text() for item in axis.get_legend().get_texts()], classes)
            self.assertEqual([item.get_text() for item in axis.texts], ["Skipped"])
            for column, bars in enumerate(axis.containers):
                np.testing.assert_allclose([bar.get_width() for bar in bars], fractions[:, column] * 100)
                np.testing.assert_allclose([bar.get_x() for bar in bars], fractions[:, :column].sum(axis=1) * 100)

        with patch.object(plt, "show", side_effect=inspect_chart) as show:
            sesca.plot_sesca_results(output)
        show.assert_called_once_with(block=True)
        self.assertFalse(plt.fignum_exists(figures[0].number))

    def test_cli_plot_opens_after_workbook_is_saved(self):
        folder, output = self.write_combined()

        def inspect_saved_results(*args, **kwargs):
            saved = load_workbook(output, read_only=True, data_only=True)
            try:
                self.assertIn("SESCA_results", saved.sheetnames)
                self.assertEqual(saved["SESCA_results"]["D2"].value, 0.2)
            finally:
                saved.close()
            bars = plt.gcf().axes[0].containers
            np.testing.assert_allclose([group[0].get_width() for group in bars], [20, 30, 50])

        arguments = [
            "deconvolution.py", "--folder", str(folder), "--path-in-mm", "1",
            "--num-of-aa", "20", "--concentration-in-mM", "0.1",
            "--sesca-dir", str(self.root), "--plot",
        ]
        with patch.object(sys, "argv", arguments), patch.object(plt, "show", side_effect=inspect_saved_results) as show:
            self.assertEqual(deconvolution.main(), 0)
        show.assert_called_once_with(block=True)

    def test_all_skipped_results_do_not_open_an_empty_plot(self):
        output = self.root / "skipped.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "SESCA_results"
        sheet.append(["spectrum", "status", "Alpha_fraction", "Beta_fraction", "Coil_fraction"])
        sheet.append(["sample", "skipped", None, None, None])
        workbook.save(output)
        workbook.close()
        stdout = io.StringIO()
        with patch.object(plt, "show") as show, patch("sys.stdout", stdout):
            sesca.plot_sesca_results(output)
        show.assert_not_called()
        self.assertIn("No fitted spectra", stdout.getvalue())

    def test_plot_with_mre_only_does_not_modify_workbook(self):
        folder, output = self.write_combined()
        original = output.read_bytes()
        with self.assertRaisesRegex(ValueError, "--mre-only"):
            deconvolution.convert_to_mre(folder, path_in_mm=1, concentration_in_M=0.0001, num_of_aa=20, plot=True)
        self.assertEqual(output.read_bytes(), original)


@unittest.skipUnless((sesca.DEFAULT_SESCA_DIR / "scripts" / "SESCA_deconv.py").is_file(), "Local official SESCA package not installed")
class RealSescaTests(unittest.TestCase):
    def setUp(self):
        self.options = sesca.SescaOptions()
        _, basis_file = sesca.resolve_sesca(self.options)
        _, self.basis = sesca.read_basis(basis_file, self.options)
        self.fractions = np.asarray([0.25, 0.35, 0.40])
        self.observed = self.basis[:, 1:] @ self.fractions * 1000

    def test_official_engine_recovers_known_mixture_and_is_reproducible(self):
        # Multiply by two to independently check SESCA's inferred intensity scaling.
        rows = [[wave, value * 2] for wave, value in zip(self.basis[:, 0], self.observed)]
        fit = sesca.fit_spectra(["known mixture"], rows, self.options).fits[0]
        self.assertEqual(fit.status, "ok")
        np.testing.assert_allclose(fit.fractions, self.fractions, atol=2e-4)
        self.assertAlmostEqual(fit.scale, 0.5, places=4)
        self.assertLess(fit.rmsd_kmre, 0.002)
        repeated = sesca.fit_spectra(["known mixture"], rows, self.options).fits[0]
        np.testing.assert_allclose(fit.fractions, repeated.fractions, rtol=0, atol=0)

    def test_default_cli_runs_real_sesca_and_writes_full_precision_fits(self):
        with tempfile.TemporaryDirectory(prefix="SESCA CLI ") as temporary:
            folder = Path(temporary).resolve() / "data with spaces"
            folder.mkdir()
            output = folder.parent / f"{folder.name}-combined.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "CD_subtr"
            sheet.append(["wavelength_nm", "peptide", "zero", "peptide"])
            for wave, value in zip(self.basis[:, 0], self.observed):
                sheet.append([wave, value * 0.002, 0, value * 0.002])
            workbook.save(output)
            workbook.close()
            result = subprocess.run(
                [sys.executable, str(PROJECT_ROOT / "deconvolution.py"), "--folder", str(folder),
                 "--path-in-mm", "1", "--num-of-aa", "20", "--concentration-in-mM", "0.1"],
                capture_output=True, text=True, check=False,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("SESCA_results", result.stdout)
            with output.open("rb") as stream:
                saved = load_workbook(stream, data_only=True)
                fractions = [saved["SESCA_results"].cell(2, col).value for col in (4, 5, 6)]
                np.testing.assert_allclose(fractions, self.fractions, atol=2e-4)
                self.assertEqual(saved["SESCA_fit"].max_row, 62)
                saved.close()


if __name__ == "__main__":
    unittest.main()
