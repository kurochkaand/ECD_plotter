"""Regression checks using temporary instrument exports and Excel workbooks."""

import os
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook

matplotlib.use("Agg")
import matplotlib.pyplot as plt

import merge
import spectral_component_analysis as analysis


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def quadratic_window_fit(values, window):
    """Independent least-squares oracle for sample-aligned smoothing."""
    fitted = []
    for row in range(len(values)):
        start = min(max(row - window // 2, 0), len(values) - window)
        offsets = np.arange(start, start + window, dtype=float) - row
        design = np.column_stack([np.ones(window), offsets, offsets**2])
        coefficients, *_ = np.linalg.lstsq(design, values[start:start + window], rcond=None)
        fitted.append(coefficients[0])
    return np.array(fitted)


class MergeTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="ECD cleanup ")
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name) / "instrument data"
        self.folder.mkdir()
        self.write_spectrum("sample", [5, 8, 11])

    def write_spectrum(
        self, name, cd, wavelengths=(200, 201, 202), absorbance=0.25,
        date="26/10/14", time="18:41:36", ht=500,
    ):
        rows = []
        ht_values = np.repeat(ht, len(cd)) if np.isscalar(ht) else ht
        abs_values = np.repeat(absorbance, len(cd)) if np.isscalar(absorbance) or absorbance is None else absorbance
        for wavelength, signal, voltage, abs_value in zip(wavelengths, cd, ht_values, abs_values):
            values = [wavelength, signal, voltage]
            if absorbance is not None:
                values.append(abs_value)
            rows.append(",".join(map(str, values)) + "\n")
        (self.folder / f"{name}.csv").write_text(
            "TITLE,Instrument export\nDATA TYPE,\nORIGIN,JASCO\nOWNER,\n"
            f"DATE,{date}\nTIME,{time}\nXYDATA\n" + "".join(rows)
            + "##### Extended Information\n",
            encoding="latin-1",
        )

    def write_coefficients(self, rows):
        (self.folder / "concentration_table.txt").write_text(
            "spectrum\tconcentration_coefficient\n" + rows,
            encoding="utf-8",
        )

    def test_workbook_preserves_raw_values_and_normalizes_without_subtraction(self):
        self.write_coefficients("sample\t2\n")
        output = merge.merge_spectra(self.folder)
        self.assertEqual(output, self.folder.parent / "instrument data-combined.xlsx")
        with pd.ExcelFile(output) as workbook:
            self.assertEqual(workbook.sheet_names, ["CD", "HT", "Abs", "CD_subtr", "Abs_subtr"])
            cd = pd.read_excel(workbook, sheet_name="CD")
            ht = pd.read_excel(workbook, sheet_name="HT")
            absorbance = pd.read_excel(workbook, sheet_name="Abs")
            summary = pd.read_excel(workbook, sheet_name="CD_subtr", header=None)
        for channel in (cd, ht, absorbance):
            self.assertEqual(channel.columns.tolist(), ["wavelength_nm", "sample"])
            self.assertEqual(channel["wavelength_nm"].tolist(), [200, 201, 202])
        self.assertEqual(cd["sample"].tolist(), [5, 8, 11])
        self.assertEqual(ht["sample"].tolist(), [500, 500, 500])
        self.assertEqual(absorbance["sample"].tolist(), [0.25, 0.25, 0.25])
        self.assertEqual(
            summary.iloc[0].tolist(),
            ["wavelength_nm", "sample", "zero", "sample"],
        )
        np.testing.assert_allclose(
            summary.iloc[1:].to_numpy(dtype=float),
            [[200, 5, 0, 2.5], [201, 8, 0, 4], [202, 11, 0, 5.5]],
        )

    def test_first_run_creates_default_concentration_table(self):
        merge.merge_spectra(self.folder)
        table = pd.read_csv(self.folder / "concentration_table.txt", sep="\t")
        self.assertEqual(table.to_dict("list"), {
            "spectrum": ["sample"], "concentration_coefficient": [1.0], "pH": [7.0],
            "short name": ["sample"],
        })

    def test_new_table_omits_reference_and_deleting_table_recreates_defaults(self):
        self.write_spectrum("H2O_reference", [1, 2, 3])
        merge.merge_spectra(self.folder, subtract="H2O")
        table_file = self.folder / "concentration_table.txt"
        self.assertEqual(pd.read_csv(table_file, sep="\t")["spectrum"].tolist(), ["sample"])
        original = table_file.read_bytes()

        self.write_spectrum("new sample", [4, 6, 8], time="19:00:00")
        output = merge.merge_spectra(self.folder, subtract="H2O")
        self.assertEqual(pd.read_excel(output, sheet_name="CD").columns.tolist(),
                         ["wavelength_nm", "H2O_reference", "sample"])
        self.assertEqual(table_file.read_bytes(), original)

        table_file.unlink()
        merge.merge_spectra(self.folder, subtract="H2O")
        table = pd.read_csv(table_file, sep="\t")
        self.assertEqual(table["spectrum"].tolist(), ["sample", "new sample"])
        self.assertEqual(table["concentration_coefficient"].tolist(), [1, 1])

    def test_edited_table_controls_all_sheets_and_plots_without_reading_excluded_csvs(self):
        self.write_spectrum("removed", [100, 200, 300])
        merge.merge_spectra(self.folder)
        table_file = self.folder / "concentration_table.txt"
        original = (
            "\ufeffspectrum\t\tconcentration_coefficient\t\tPh\t\tshort name\t\r\n"
            "sample\t\t2\t\t6.8\t\tMy label\t\r\n\r\n"
        ).encode("utf-8")
        table_file.write_bytes(original)
        # Neither an excluded old spectrum nor a new CSV needs valid metadata/data.
        for name in ("removed", "new spectrum"):
            (self.folder / f"{name}.csv").write_text("ignored malformed CSV", encoding="utf-8")

        def inspect_plot(*args, **kwargs):
            for axis in plt.gcf().axes:
                self.assertEqual([line.get_label() for line in axis.lines], ["My label"])
            np.testing.assert_allclose(plt.gcf().axes[0].lines[0].get_ydata(), [2.5, 4, 5.5])

        with patch.object(plt, "show", side_effect=inspect_plot), \
                patch.object(merge, "read_spectrum", wraps=merge.read_spectrum) as read:
            output = merge.merge_spectra(self.folder, plot=True, smooth=3, ht_threshold=500)
        self.assertEqual([call.args[0].stem for call in read.call_args_list], ["sample"])
        workbook = load_workbook(output)
        try:
            for sheet in ("CD", "HT", "Abs"):
                self.assertEqual([cell.value for cell in workbook[sheet][1]],
                                 ["wavelength_nm", "My label"])
            for sheet in ("CD_subtr", "Abs_subtr"):
                self.assertEqual([cell.value for cell in workbook[sheet][1]],
                                 ["wavelength_nm", "My label", "zero", "My label"])
        finally:
            workbook.close()
        self.assertEqual(table_file.read_bytes(), original)

    def test_short_names_remove_common_affixes_and_keep_unique_nonempty_labels(self):
        for names, expected in (
            (["run_first_export", "run_second_export"], ["first", "second"]),
            (["first_export", "second_export"], ["first", "second"]),
            (["run_first", "run_second"], ["first", "second"]),
            (["only_sample_export"], ["only_sample_export"]),
            (["sample", "later sample"], ["sample", "later sample"]),
            (["run_A_export", "run__A_export", "run_B_export"],
             ["run_A_export", "run__A_export", "run_B_export"]),
            (["run_zero_export", "run_sample_export"],
             ["run_zero_export", "run_sample_export"]),
        ):
            with self.subTest(names=names):
                labels = merge.determine_short_names(names)
                self.assertEqual([labels[name] for name in names], expected)
        names = ["run_A_export", "run_B_export"]
        self.assertEqual(merge.determine_short_names(names, subtraction_reference="A"),
                         dict(zip(names, names)))

    def test_short_names_ignore_subtraction_reference_in_table_workbook_and_plot(self):
        first = "peptide_pH7.43_1mm_185-300nm_4sc_20nm-1"
        second = "peptide_pH5.77_1mm_185-300nm_4sc_20nm-1"
        reference = "H2O_1mm_185-300nm_4sc_20nm-1"
        (self.folder / "sample.csv").rename(self.folder / f"{first}.csv")
        self.write_spectrum(second, [4, 6, 8], time="19:00:00")
        self.write_spectrum(reference, [1, 2, 3], time="20:00:00")
        table_file = self.folder / "concentration_table.txt"
        table_file.write_text(
            "spectrum\tconcentration_coefficient\tPh\n"
            f"{second}\t4\t6.8\n{reference}\t1\t7\n{first}\t2\t6.5\n",
            encoding="utf-8",
        )
        contents = table_file.read_bytes()

        def inspect_plot(*args, **kwargs):
            for axis in plt.gcf().axes:
                self.assertEqual([line.get_label() for line in axis.lines], ["7.43", "5.77"])
            np.testing.assert_allclose(plt.gcf().axes[0].lines[0].get_ydata(), [2, 3, 4])

        with patch.object(plt, "show", side_effect=inspect_plot):
            output = merge.merge_spectra(self.folder, subtract="H2O", plot=True)
        workbook = load_workbook(output)
        try:
            raw_names = ["wavelength_nm", "7.43", "5.77", reference]
            for sheet in ("CD", "HT", "Abs"):
                self.assertEqual([cell.value for cell in workbook[sheet][1]], raw_names)
            for sheet in ("CD_subtr", "Abs_subtr"):
                self.assertEqual([cell.value for cell in workbook[sheet][1]],
                                 [*raw_names, "zero", "7.43", "5.77"])
                self.assertEqual(workbook[sheet]["F2"].value, "=(B2-$D2)/2.0")
        finally:
            workbook.close()
        table = merge.load_concentration_table(
            self.folder, [first, second, reference], subtraction_reference=reference,
        )
        self.assertEqual(table.index.tolist(), [first, second])
        self.assertEqual(table["short name"].tolist(), ["7.43", "5.77"])
        self.assertEqual(table["concentration_coefficient"].tolist(), [2, 4])
        self.assertEqual(table["pH"].astype(float).tolist(), [6.5, 6.8])
        merge.merge_spectra(self.folder, subtract="H2O")
        self.assertEqual(table_file.read_bytes(), contents)

    def test_four_column_table_preserves_short_names_and_only_infers_missing_ph_in_memory(self):
        names = ["peptide_pH7.43_export", "peptide_pH5.77_export"]
        for aligned in (False, True):
            with self.subTest(aligned=aligned):
                separator = "\t\t" if aligned else "\t"
                table_file = self.folder / "concentration_table.txt"
                table_file.write_text(
                    separator.join(merge.CONCENTRATION_COLUMNS) + "\n"
                    + separator.join([names[1], "3", "6.5", "outdated"]) + "\n"
                    + f"{names[0]}\t2\t\told label\n",
                    encoding="utf-8",
                )
                contents = table_file.read_bytes()
                table = merge.load_concentration_table(self.folder, names)
                self.assertEqual(table["short name"].tolist(), ["old label", "outdated"])
                self.assertEqual(table["pH"].astype(float).tolist(), [7.43, 6.5])
                self.assertEqual(table["concentration_coefficient"].tolist(), [2, 3])
                self.assertEqual(table_file.read_bytes(), contents)

    def test_ph_is_extracted_from_filename_tokens(self):
        for name, expected in (
            ("Cd-peptide_0.2ekv_pH7.43_1mm_185-300nm_4sc_20nm-1", 7.43),
            ("Cd-peptide_1.6ekv_pH5.77_1mm_185-300nm_4sc_20nm-1", 5.77),
            ("peptide_pH2.18_201uL_1mm_185-300nm_4sc_20nm-1", 2.18),
            ("peptide_pH7.9_0uL_1mm_185-300nm_4sc_20nm-1", 7.9),
            ("peptide_pH 7.01_21uL_1mm_185-300nm_4sc_20nm-1", 7.01),
            ("peptide_PH8_1mm", 8.0),
            ("pH0", 0.0),
            ("peptide_pH-0.5_1mm", -0.5),
            ("H2O_1mm_185-300nm_4sc_20nm-1", None),
            ("phosphate_7.43_1mm", None),
        ):
            with self.subTest(name=name):
                self.assertEqual(merge.extract_ph_from_name(name), expected)

    def test_new_table_uses_filename_ph_and_default_for_missing_tokens(self):
        names = [
            "Cd-peptide_0.2ekv_pH7.43_1mm_185-300nm_4sc_20nm-1",
            "peptide_pH7.9_0uL_1mm_185-300nm_4sc_20nm-1",
            "H2O_1mm_185-300nm_4sc_20nm-1",
        ]
        merge.load_concentration_coefficients(self.folder, names)
        table = pd.read_csv(self.folder / "concentration_table.txt", sep="\t")
        self.assertEqual(table["spectrum"].tolist(), names)
        self.assertEqual(table["pH"].tolist(), [7.43, 7.9, 7.0])
        self.assertEqual(table["concentration_coefficient"].tolist(), [1.0, 1.0, 1.0])

    def test_missing_ph_in_existing_tables_is_inferred(self):
        for contents in (
            "spectrum\tconcentration_coefficient\nsample_pH5.77\t2\n",
            "spectrum\tconcentration_coefficient\tPh\nsample_pH5.77\t2\t\n",
        ):
            with self.subTest(contents=contents):
                table_file = self.folder / "concentration_table.txt"
                table_file.write_text(contents, encoding="utf-8")
                original = table_file.read_bytes()
                table = merge.load_concentration_table(self.folder, ["sample_pH5.77"])
                self.assertEqual(table["pH"].tolist(), [5.77])
                self.assertEqual(table["concentration_coefficient"].tolist(), [2])
                self.assertEqual(table_file.read_bytes(), original)

    def test_existing_ph_is_preserved(self):
        table_file = self.folder / "concentration_table.txt"
        table_file.write_text(
            "spectrum\tconcentration_coefficient\tPh\n"
            "sample_pH7.43\t2\t6.8\nH2O_1mm\t3\t6.5\n",
            encoding="utf-8",
        )
        names = ["sample_pH7.43", "H2O_1mm"]
        original = table_file.read_bytes()
        table = merge.load_concentration_table(self.folder, names)
        self.assertEqual(table["pH"].astype(float).tolist(), [6.8, 6.5])
        self.assertEqual(table["concentration_coefficient"].tolist(), [2, 3])
        self.assertEqual(table_file.read_bytes(), original)

    def test_concentration_table_accepts_alignment_tabs(self):
        names = ["H2O reference", "sample_pH5.77"]
        for contents, expected_ph in (
            (
                "spectrum\t\tconcentration_coefficient\t\n"
                "H2O reference\t\t\t\t3\t\n"
                " sample_pH5.77 \t\t2\n",
                [7.0, 5.77],
            ),
            (
                "spectrum\tconcentration_coefficient\tPh\n"
                "H2O reference\t\t\t\t3\t7.0\n"
                "sample_pH5.77\t\t2\t\t6.8\t\n",
                [7.0, 6.8],
            ),
            (
                "spectrum\t\tconcentration_coefficient\t\tPh\t\n"
                "\n\t\t\n  \t  \t\n"
                "H2O reference\t\t3\t\t\n"
                "sample_pH5.77\t2\t\n",
                [7.0, 5.77],
            ),
            (
                "spectrum\t\tconcentration_coefficient\t\tPh\tshort name\t\n"
                "H2O reference\t\t3\t7.0\n"
                "sample_pH5.77\t2\t6.8\n",
                [7.0, 6.8],
            ),
        ):
            with self.subTest(contents=contents):
                table_file = self.folder / "concentration_table.txt"
                table_file.write_text(contents, encoding="utf-8-sig")
                original = table_file.read_bytes()
                table = merge.load_concentration_table(self.folder, names)
                self.assertEqual(table["concentration_coefficient"].to_dict(), dict(zip(names, [3, 2])))
                self.assertEqual(table.index.tolist(), names)
                self.assertEqual(table["pH"].astype(float).tolist(), expected_ph)
                self.assertEqual(table_file.read_bytes(), original)

    def test_duplicate_spectra_with_alignment_tabs_report_names(self):
        self.write_coefficients("sample\t\t\t1\nsample\t\t2\n")
        with self.assertRaisesRegex(ValueError, "Duplicate spectra.*: sample$"):
            merge.load_concentration_coefficients(self.folder, ["sample"])

    def test_concentration_table_rejects_extra_values_without_rewriting(self):
        table_file = self.folder / "concentration_table.txt"
        contents = (
            "spectrum\tconcentration_coefficient\tPh\n"
            "sample\t\t2\t7.0\t8.0\n"
        )
        table_file.write_text(contents, encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "tab-separated columns"):
            merge.load_concentration_coefficients(self.folder, ["sample"])
        self.assertEqual(table_file.read_text(encoding="utf-8"), contents)

    def test_cli_no_longer_offers_refresh_ph(self):
        result = subprocess.run(
            [sys.executable, str(PROJECT_ROOT / "merge.py"), "--help"],
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertNotIn("--refresh-ph", result.stdout)

    def test_workbook_follows_acquisition_datetime_without_reordering_existing_table(self):
        self.write_spectrum("a_later", [4, 8, 12], date="26/10/15", time="00:00:00")
        self.write_spectrum("z_earlier", [5, 10, 15], date="26/10/13", time="11:59:59 PM")
        self.write_spectrum("morning", [3, 6, 9], time="6:41:36 AM")
        table_file = self.folder / "concentration_table.txt"
        table_file.write_text(
            "spectrum\t\t\t\tconcentration_coefficient\tPh\n"
            "a_later\t4\t6.5\nsample\t2\t7.2\n"
            "z_earlier\t5\t8.1\nmorning\t3\t6.8\n",
            encoding="utf-8",
        )
        original = table_file.read_bytes()
        expected_names = ["z_earlier", "morning", "sample", "a_later"]

        output = merge.merge_spectra(self.folder)
        with pd.ExcelFile(output) as workbook:
            for sheet_name in ("CD", "HT", "Abs"):
                channel = pd.read_excel(workbook, sheet_name=sheet_name)
                self.assertEqual(channel.columns.tolist(), ["wavelength_nm", *expected_names])
        summary = pd.read_excel(output, sheet_name="CD_subtr", header=None)
        self.assertEqual(
            summary.iloc[0].tolist(),
            ["wavelength_nm", *expected_names, "zero", *expected_names],
        )
        self.assertEqual(table_file.read_bytes(), original)
        np.testing.assert_allclose(
            summary.iloc[1:, -4:].to_numpy(dtype=float),
            [[1, 1, 2.5, 1], [2, 2, 4, 2], [3, 3, 5.5, 3]],
        )

    def test_new_concentration_table_follows_acquisition_datetime(self):
        self.write_spectrum("z_earlier", [1, 2, 3], time="06:00:00")
        merge.merge_spectra(self.folder)
        table = pd.read_csv(self.folder / "concentration_table.txt", sep="\t")
        self.assertEqual(table["spectrum"].tolist(), ["z_earlier", "sample"])
        self.assertEqual(table["concentration_coefficient"].tolist(), [1.0, 1.0])

    def test_equal_timestamps_use_filename_order(self):
        self.write_spectrum("zzz", [1, 2, 3])
        self.write_spectrum("Alpha", [1, 2, 3], time="6:41:36 PM")
        files = merge.find_spectral_files(self.folder)
        self.assertEqual([file.stem for file in files], ["Alpha", "sample", "zzz"])

    def test_acquisition_datetime_supports_24_hour_and_am_pm_times(self):
        for date, time, hour in (
            ("26/10/14", "18:41:36", 18),
            ("2026/10/14", "6:41:36 PM", 18),
            ("26/10/14", "12:41:36 AM", 0),
            ("26/10/14", "12:41:36 PM", 12),
        ):
            with self.subTest(date=date, time=time):
                self.write_spectrum("sample", [1, 2, 3], date=date, time=time)
                self.assertEqual(
                    merge.read_acquisition_datetime(self.folder / "sample.csv"),
                    datetime(2026, 10, 14, hour, 41, 36),
                )

    def test_invalid_acquisition_datetime_reports_file_and_cells(self):
        file = self.folder / "sample.csv"
        for date, time in (("", "18:41:36"), ("26/13/14", "18:41:36"), ("26/10/14", "25:00:00")):
            with self.subTest(date=date, time=time):
                self.write_spectrum("sample", [1, 2, 3], date=date, time=time)
                with self.assertRaisesRegex(ValueError, "Invalid acquisition date/time") as raised:
                    merge.find_spectral_files(self.folder)
                message = str(raised.exception)
                self.assertIn(str(file), message)
                self.assertIn(f"B5 (CSV line 5): {date!r}", message)
                self.assertIn(f"B6 (CSV line 6): {time!r}", message)

    def test_mismatched_wavelengths_do_not_write_workbook(self):
        self.write_spectrum("other sample", [5, 8, 11], wavelengths=(200, 201, 203))
        reference_file = self.folder / "sample.csv"
        actual_file = self.folder / "other sample.csv"
        with patch.object(merge, "find_spectral_files", return_value=[reference_file, actual_file]):
            with self.assertRaisesRegex(ValueError, "X axis") as raised:
                merge.merge_spectra(self.folder)
        message = str(raised.exception)
        self.assertIn(f"X axis mismatch in {actual_file}", message)
        self.assertIn(f"Reference file: {reference_file}", message)
        self.assertIn("Spectral data row: 3", message)
        self.assertIn("Expected: 202 nm (CSV line 10)", message)
        self.assertIn("Actual:   203 nm (CSV line 10)", message)
        self.assertIn("Data rows: expected 3, actual 3", message)
        self.assertFalse((self.folder.parent / "instrument data-combined.xlsx").exists())

    def test_mismatch_line_numbers_include_blank_lines(self):
        reference_file = self.folder / "sample.csv"
        actual_file = self.folder / "other sample.csv"
        actual_file.write_text(
            "Instrument export\nXYDATA\n\n200,5,500\n \n201,8,500\n"
            "203,11,500\n##### Extended Information\n",
            encoding="latin-1",
        )
        with self.assertRaises(ValueError) as raised:
            merge.validate_wavelength_axis(
                merge.read_spectrum(reference_file), merge.read_spectrum(actual_file),
                reference_file, actual_file,
            )
        self.assertIn("Expected: 202 nm (CSV line 10)", str(raised.exception))
        self.assertIn("Actual:   203 nm (CSV line 7)", str(raised.exception))

    def test_different_axis_lengths_report_the_first_missing_or_extra_row(self):
        reference_file = self.folder / "sample.csv"
        actual_file = self.folder / "other sample.csv"
        for wavelengths, row, expected, actual in (
            ((200, 201), 3, "202 nm (CSV line 10)",
             "<end of spectral data> (CSV line 10)"),
            ((200, 201, 202, 203), 4, "<end of spectral data> (CSV line 11)",
             "203 nm (CSV line 11)"),
        ):
            with self.subTest(wavelengths=wavelengths):
                self.write_spectrum("other sample", range(len(wavelengths)), wavelengths)
                with self.assertRaises(ValueError) as raised:
                    merge.validate_wavelength_axis(
                        merge.read_spectrum(reference_file), merge.read_spectrum(actual_file),
                        reference_file, actual_file,
                    )
                message = str(raised.exception)
                self.assertIn(f"Spectral data row: {row}", message)
                self.assertIn(f"Expected: {expected}", message)
                self.assertIn(f"Actual:   {actual}", message)
                self.assertIn(f"Data rows: expected 3, actual {len(wavelengths)}", message)

    def test_equivalent_integer_and_decimal_wavelengths_are_accepted(self):
        self.write_spectrum("other sample", [1, 2, 3], wavelengths=(200.0, 201.0, 202.0))
        output = merge.merge_spectra(self.folder)
        self.assertTrue(output.is_file())

    def test_four_column_exports_are_written_to_separate_channel_sheets(self):
        self.write_spectrum("sample", [5, 8, 11], absorbance=0.25)
        self.write_spectrum("other sample", [1, 2, 3], absorbance=0.5, ht=650)
        frame = merge.read_spectrum(self.folder / "sample.csv")
        self.assertEqual(frame["wavelength_nm"].tolist(), [200, 201, 202])
        self.assertEqual(frame["cd_mdeg"].tolist(), [5, 8, 11])
        self.assertEqual(frame["ht_v"].tolist(), [500, 500, 500])
        self.assertEqual(frame["absorbance"].tolist(), [0.25, 0.25, 0.25])
        self.assertEqual(frame.index.tolist(), [0, 1, 2])
        output = merge.merge_spectra(self.folder)
        with pd.ExcelFile(output) as workbook:
            self.assertEqual(workbook.sheet_names, ["CD", "HT", "Abs", "CD_subtr", "Abs_subtr"])
            for sheet_name, expected_sample, expected_other in (
                ("CD", [5, 8, 11], [1, 2, 3]),
                ("HT", [500, 500, 500], [650, 650, 650]),
                ("Abs", [0.25, 0.25, 0.25], [0.5, 0.5, 0.5]),
            ):
                with self.subTest(sheet_name=sheet_name):
                    channel = pd.read_excel(workbook, sheet_name=sheet_name)
                    self.assertEqual(channel["wavelength_nm"].tolist(), [200, 201, 202])
                    self.assertEqual(channel["sample"].tolist(), expected_sample)
                    self.assertEqual(channel["other sample"].tolist(), expected_other)

    def test_three_column_exports_leave_absorbance_blank(self):
        self.write_spectrum("sample", [5, 8, 11], absorbance=None)
        output = merge.merge_spectra(self.folder)
        absorbance = pd.read_excel(output, sheet_name="Abs")
        self.assertEqual(absorbance["wavelength_nm"].tolist(), [200, 201, 202])
        self.assertTrue(absorbance["sample"].isna().all())

    def test_absorbance_normalization_preserves_measurements_with_cd_processing(self):
        wavelengths = np.arange(200, 205)
        values = [0.4, 1.6, 0.8, 2.4, 1.2]
        self.write_spectrum("sample", [5, 8, 11, 7, 3], wavelengths=wavelengths, absorbance=values, ht=501)
        self.write_coefficients("sample\t2\n")

        def inspect_chart(*args, **kwargs):
            cd_axis, abs_axis = plt.gcf().axes
            self.assertTrue(np.isnan(cd_axis.get_lines()[0].get_ydata()).all())
            self.assertNotIn("HT cutoff", abs_axis.get_title())
            self.assertNotIn("Savitzky-Golay", abs_axis.get_title())
            np.testing.assert_allclose(abs_axis.get_lines()[0].get_ydata(), [0.2, 0.8, 0.4, 1.2, 0.6])

        with patch.object(plt, "show", side_effect=inspect_chart):
            output = merge.merge_spectra(self.folder, smooth=5, ht_threshold=500, plot=True)
        workbook = load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)
        sheet = workbook["Abs_subtr"]
        self.assertEqual([cell.value for cell in sheet[1]], ["wavelength_nm", "sample", "zero", "sample"])
        self.assertEqual([sheet.cell(row, 1).value for row in range(2, 7)], wavelengths.tolist())
        for sheet_name in ("Abs", "Abs_subtr"):
            self.assertEqual([workbook[sheet_name].cell(row, 2).value for row in range(2, 7)], values)
        self.assertEqual([sheet.cell(row, 3).value for row in range(2, 7)], [0] * 5)
        np.testing.assert_allclose([sheet.cell(row, 4).value for row in range(2, 7)], [0.2, 0.8, 0.4, 1.2, 0.6])

    def test_absorbance_subtraction_matches_sheet_formulas_and_plot(self):
        self.write_spectrum("sample", [5, 8, 11], absorbance=[0.4, 0.8, 1.2])
        self.write_spectrum("H2O_reference", [1, 2, 3], absorbance=[0.1, 0.2, 0.3], time="19:00:00")
        self.write_spectrum("later sample", [4, 6, 8], absorbance=[0.05, 0.4, 0.9], time="20:00:00")
        self.write_coefficients("sample\t2\nH2O_reference\t3\nlater sample\t1\n")

        def inspect_chart(*args, **kwargs):
            cd_axis, abs_axis = plt.gcf().axes
            self.assertEqual(abs_axis.get_xlabel(), "Wavelength (nm)")
            self.assertEqual(abs_axis.get_ylabel(), "Normalized absorbance (AU)")
            self.assertIn("H2O_reference", abs_axis.get_title())
            lines = abs_axis.get_lines()
            self.assertEqual([line.get_label() for line in lines], ["sample", "later sample"])
            np.testing.assert_allclose(lines[0].get_xdata(), [200, 201, 202])
            np.testing.assert_allclose(lines[0].get_ydata(), [0.15, 0.3, 0.45])
            np.testing.assert_allclose(lines[1].get_ydata(), [-0.05, 0.2, 0.6])
            for cd_line, abs_line in zip(cd_axis.get_lines(), lines):
                np.testing.assert_allclose(cd_line.get_color(), abs_line.get_color())

        with patch.object(plt, "show", side_effect=inspect_chart) as show:
            output = merge.merge_spectra(self.folder, subtract="h2o", plot=True)
        show.assert_called_once_with(block=True)
        workbook = load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)
        sheet = workbook["Abs_subtr"]
        self.assertEqual([cell.value for cell in sheet[1]], [
            "wavelength_nm", "sample", "H2O_reference", "later sample", "zero", "sample", "later sample",
        ])
        for row in range(2, 5):
            self.assertEqual(sheet.cell(row, 6).value, f"=(B{row}-$C{row})/2.0")
            self.assertEqual(sheet.cell(row, 7).value, f"=D{row}-$C{row}")
            self.assertEqual(sheet.cell(row, 6).data_type, "f")
            self.assertEqual(sheet.cell(row, 7).data_type, "f")
        for column, expected in ((2, [0.4, 0.8, 1.2]), (3, [0.1, 0.2, 0.3]), (4, [0.05, 0.4, 0.9])):
            for sheet_name in ("Abs", "Abs_subtr"):
                self.assertEqual([workbook[sheet_name].cell(row, column).value for row in range(2, 5)], expected)

    def test_missing_absorbance_stays_blank_in_corrections_and_plot(self):
        for reference_abs in (None, [0.1, np.nan, 0.3]):
            with self.subTest(reference_abs=reference_abs):
                self.write_spectrum("sample", [5, 8, 11], absorbance=[0.4, 0.8, np.nan])
                self.write_spectrum("H2O_reference", [1, 2, 3], absorbance=reference_abs, time="19:00:00")
                self.write_spectrum("legacy sample", [4, 6, 8], absorbance=None, time="20:00:00")
                self.write_coefficients("sample\t2\nH2O_reference\t3\nlegacy sample\t1\n")

                def inspect_chart(*args, **kwargs):
                    axes = plt.gcf().axes
                    self.assertEqual(len(axes), 1 if reference_abs is None else 2)
                    if reference_abs is not None:
                        lines = axes[1].get_lines()
                        self.assertEqual([line.get_label() for line in lines], ["sample"])
                        np.testing.assert_allclose(lines[0].get_ydata(), [0.15, np.nan, np.nan], equal_nan=True)

                with patch.object(plt, "show", side_effect=inspect_chart):
                    output = merge.merge_spectra(self.folder, subtract="H2O", plot=True)
                workbook = load_workbook(output, data_only=False)
                try:
                    sheet = workbook["Abs_subtr"]
                    self.assertEqual(sheet["F2"].value, None if reference_abs is None else "=(B2-$C2)/2.0")
                    for coordinate in ("F3", "F4", "G2", "G3", "G4"):
                        self.assertIsNone(sheet[coordinate].value)
                        self.assertNotEqual(sheet[coordinate].data_type, "f")
                finally:
                    workbook.close()

    def test_all_filenames_are_treated_as_samples(self):
        self.write_spectrum("water", [1, 2, 3])
        self.write_spectrum("Water blank", [2, 4, 6])
        self.write_coefficients("sample\t2\nwater\t4\nWater blank\t2\n")
        self.assertEqual(len(merge.find_spectral_files(self.folder)), 3)

        output = merge.merge_spectra(self.folder)
        summary = pd.read_excel(output, sheet_name="CD_subtr", header=None)
        headers = summary.iloc[0].tolist()
        for name, expected in (
            ("sample", [2.5, 4, 5.5]),
            ("water", [0.25, 0.5, 0.75]),
            ("Water blank", [1, 2, 3]),
        ):
            with self.subTest(name=name):
                columns = [index for index, header in enumerate(headers) if header == name]
                self.assertEqual(len(columns), 2)
                np.testing.assert_allclose(
                    summary.iloc[1:, columns[-1]].to_numpy(dtype=float), expected
                )

    def test_subtraction_writes_formulas_for_each_wavelength_and_preserves_raw_data(self):
        self.write_spectrum("H2O_reference", [1, 2, 3], time="19:00:00")
        self.write_spectrum("later sample", [4, 6, 8], time="20:00:00")
        self.write_coefficients("sample\t2\nlater sample\t1\n")
        output = merge.merge_spectra(self.folder, subtract="h2o")
        workbook = load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["CD", "HT", "Abs", "CD_subtr", "Abs_subtr"])
        sheet = workbook["CD_subtr"]
        self.assertEqual([cell.value for cell in sheet[1]], [
            "wavelength_nm", "sample", "H2O_reference", "later sample", "zero",
            "sample", "later sample",
        ])
        for row in range(2, 5):
            self.assertEqual(sheet.cell(row, 6).value, f"=(B{row}-$C{row})/2.0")
            self.assertEqual(sheet.cell(row, 7).value, f"=D{row}-$C{row}")
            self.assertEqual(sheet.cell(row, 6).data_type, "f")
            self.assertEqual(sheet.cell(row, 7).data_type, "f")
        self.assertEqual([sheet.cell(row, 2).value for row in range(2, 5)], [5, 8, 11])
        self.assertEqual([sheet.cell(row, 3).value for row in range(2, 5)], [1, 2, 3])
        self.assertEqual(workbook["CD"]["B2"].value, 5)
        self.assertEqual(workbook["CD"]["C2"].value, 1)
        self.assertEqual(workbook["HT"]["B2"].value, 500)
        self.assertEqual(workbook["Abs"]["B2"].value, 0.25)
        self.assertEqual(workbook.calculation.calcMode, "auto")
        self.assertTrue(workbook.calculation.fullCalcOnLoad)
        self.assertTrue(workbook.calculation.forceFullCalc)

    def test_subtraction_handles_reference_columns_after_z(self):
        for index in range(25):
            self.write_spectrum(f"other_{index:02}", [4, 6, 8], time="19:00:00")
        self.write_spectrum("H2O_reference", [1, 2, 3], time="20:00:00")
        output = merge.merge_spectra(self.folder, subtract="H2O")
        workbook = load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)
        sheet = workbook["CD_subtr"]
        self.assertEqual(sheet["AB1"].value, "H2O_reference")
        self.assertEqual(sheet["AD2"].value, "=B2-$AB2")
        self.assertEqual(sheet["AD4"].value, "=B4-$AB4")
        self.assertEqual(workbook["Abs_subtr"]["AD2"].value, "=B2-$AB2")
        self.assertEqual(workbook["Abs_subtr"]["AD4"].value, "=B4-$AB4")

    def test_missing_or_ambiguous_subtraction_reference_does_not_write_outputs(self):
        table_file = self.folder / "concentration_table.txt"
        for prefix, expected in (("H2O", "found 0"), (" ", "cannot be empty")):
            with self.subTest(prefix=prefix):
                with self.assertRaisesRegex(ValueError, expected):
                    merge.merge_spectra(self.folder, subtract=prefix)
                self.assertFalse(table_file.exists())
                self.assertFalse((self.folder.parent / "instrument data-combined.xlsx").exists())

        self.write_spectrum("H2O_1", [1, 2, 3])
        self.write_spectrum("H2O_2", [2, 3, 4])
        with self.assertRaisesRegex(ValueError, "found 2") as raised:
            merge.merge_spectra(self.folder, subtract="H2O")
        self.assertIn("H2O_1", str(raised.exception))
        self.assertIn("H2O_2", str(raised.exception))
        self.assertFalse(table_file.exists())

    def test_cli_subtract_option_writes_excel_formulas(self):
        self.write_spectrum("H2O_reference", [1, 2, 3], time="19:00:00")
        result = subprocess.run(
            [sys.executable, str(PROJECT_ROOT / "merge.py"),
             "--folder", str(self.folder), "--subtract", "H2O"],
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        workbook = load_workbook(
            self.folder.parent / "instrument data-combined.xlsx", data_only=False
        )
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["CD_subtr"]["E2"].value, "=B2-$C2")
        self.assertEqual(workbook["Abs_subtr"]["E2"].value, "=B2-$C2")

    def test_plot_is_disabled_by_default(self):
        with patch.object(merge, "plot_corrected_spectra") as plot:
            merge.merge_spectra(self.folder)
        plot.assert_not_called()

    def test_plot_displays_normalized_cd_in_acquisition_order(self):
        self.write_spectrum("a later sample", [-2, 4, 12], time="19:00:00")
        self.write_coefficients("sample\t2\na later sample\t4\n")
        figures = []

        def inspect_chart(*args, **kwargs):
            figure = plt.gcf()
            figures.append(figure)
            self.assertEqual(len(figure.axes), 2)
            axis = figure.axes[0]
            self.assertEqual(axis.get_xlabel(), "Wavelength (nm)")
            self.assertEqual(axis.get_ylabel(), "Normalized CD (mdeg)")
            lines = axis.get_lines()
            self.assertEqual([line.get_label() for line in lines], ["sample", "a later sample"])
            np.testing.assert_allclose(lines[0].get_xdata(), [200, 201, 202])
            np.testing.assert_allclose(lines[0].get_ydata(), [2.5, 4, 5.5])
            np.testing.assert_allclose(lines[1].get_ydata(), [-0.5, 1, 3])
            abs_axis = figure.axes[1]
            self.assertEqual(abs_axis.get_ylabel(), "Normalized absorbance (AU)")
            abs_lines = abs_axis.get_lines()
            self.assertEqual([line.get_label() for line in abs_lines], ["sample", "a later sample"])
            np.testing.assert_allclose(abs_lines[0].get_ydata(), [0.125] * 3)
            np.testing.assert_allclose(abs_lines[1].get_ydata(), [0.0625] * 3)
            self.assertTrue((self.folder.parent / "instrument data-combined.xlsx").is_file())

        with patch.object(plt, "show", side_effect=inspect_chart) as show:
            merge.merge_spectra(self.folder, plot=True)
        show.assert_called_once_with(block=True)
        self.assertFalse(plt.fignum_exists(figures[0].number))

    def test_subtraction_plot_matches_corrections_and_keeps_excel_formulas(self):
        self.write_spectrum("H2O_reference", [1, 2, 3], time="19:00:00")
        self.write_spectrum("later sample", [4, 6, 8], time="20:00:00")
        self.write_coefficients("sample\t2\nH2O_reference\t3\nlater sample\t1\n")

        def inspect_chart(*args, **kwargs):
            axis = plt.gcf().axes[0]
            self.assertIn("H2O_reference", axis.get_title())
            lines = axis.get_lines()
            self.assertEqual([line.get_label() for line in lines], ["sample", "later sample"])
            np.testing.assert_allclose(lines[0].get_ydata(), [2, 3, 4])
            np.testing.assert_allclose(lines[1].get_ydata(), [3, 4, 5])

        with patch.object(plt, "show", side_effect=inspect_chart) as show:
            output = merge.merge_spectra(self.folder, subtract="H2O", plot=True)
        show.assert_called_once_with(block=True)
        workbook = load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["CD_subtr"]["F2"].value, "=(B2-$C2)/2.0")
        self.assertEqual(workbook["CD_subtr"]["G2"].value, "=D2-$C2")

    def test_cli_plot_option_runs_with_subtraction(self):
        self.write_spectrum("H2O_reference", [1, 2, 3], time="19:00:00")
        result = subprocess.run(
            [sys.executable, str(PROJECT_ROOT / "merge.py"),
             "--folder", str(self.folder), "--subtract", "H2O", "--plot"],
            env={**os.environ, "MPLBACKEND": "Agg"},
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Saved ", result.stdout)

    def test_even_window_smoothing_preserves_quadratics_at_original_positions(self):
        positions = np.arange(101, dtype=float)
        signal = pd.Series(
            0.005 * positions**2 + 0.3 * positions - 2,
            index=np.arange(400, 501), name="quadratic CD",
        )
        original = signal.copy()
        result = merge.smooth_cd_spectra({"quadratic": signal}, 50)["quadratic"]
        np.testing.assert_allclose(result, original, rtol=1e-12, atol=1e-12)
        self.assertEqual(result.name, signal.name)
        self.assertTrue(result.index.equals(signal.index))
        pd.testing.assert_series_equal(signal, original)

    def test_smoothing_matches_known_five_point_impulse_response(self):
        values = np.zeros(21)
        values[10] = 1
        result = merge.smooth_cd_spectra({"impulse": pd.Series(values)}, 5)["impulse"]
        expected = np.zeros(21)
        expected[8:13] = np.array([-3, 12, 17, 12, -3]) / 35
        np.testing.assert_allclose(result, expected, atol=1e-14)

    def test_fifty_point_smoothing_matches_independent_fits_including_edges(self):
        values = np.random.default_rng(2026).normal(size=101)
        result = merge.smooth_cd_spectra({"noise": pd.Series(values)}, 50)["noise"]
        np.testing.assert_allclose(result, quadratic_window_fit(values, 50), atol=1e-12)

    def test_smoothing_precedes_subtraction_in_workbook_and_plot(self):
        positions = np.arange(71, dtype=float)
        wavelengths = positions + 200
        sample = np.sin(positions / 10) + 0.1 * (-1.0)**positions
        reference = 0.1 * np.cos(positions / 8) + 0.05 * (-1.0)**positions
        self.write_spectrum("sample", sample, wavelengths=wavelengths)
        self.write_spectrum(
            "H2O_reference", reference, wavelengths=wavelengths,
            time="19:00:00", ht=650, absorbance=0.5,
        )
        self.write_coefficients("sample\t2\nH2O_reference\t3\n")
        expected_sample = quadratic_window_fit(sample, 6)
        expected_reference = quadratic_window_fit(reference, 6)

        def inspect_chart(*args, **kwargs):
            axis = plt.gcf().axes[0]
            self.assertIn("6 points, polynomial order 2", axis.get_title())
            lines = axis.get_lines()
            self.assertEqual([line.get_label() for line in lines], ["sample"])
            np.testing.assert_allclose(lines[0].get_xdata(), wavelengths)
            np.testing.assert_allclose(
                lines[0].get_ydata(), (expected_sample - expected_reference) / 2,
                atol=1e-12,
            )

        with patch.object(plt, "show", side_effect=inspect_chart) as show:
            output = merge.merge_spectra(self.folder, subtract="H2O", smooth=6, plot=True)
        show.assert_called_once_with(block=True)
        workbook = load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)
        for sheet_name, column, expected in (
            ("CD", 2, sample), ("CD", 3, reference),
            ("CD_subtr", 2, expected_sample), ("CD_subtr", 3, expected_reference),
        ):
            actual = [workbook[sheet_name].cell(row, column).value for row in range(2, 73)]
            np.testing.assert_allclose(actual, expected, atol=1e-12)
        self.assertEqual(workbook["HT"]["C2"].value, 650)
        self.assertEqual(workbook["Abs"]["C2"].value, 0.5)
        for row in range(2, 73):
            self.assertEqual(workbook["CD_subtr"].cell(row, 5).value, f"=(B{row}-$C{row})/2.0")

    def test_cli_smooth_supports_default_fifty_and_an_explicit_window(self):
        positions = np.arange(70, dtype=float)
        wavelengths = positions + 200
        values = np.sin(positions / 9) + 0.2 * (-1.0)**positions
        self.write_spectrum("sample", values, wavelengths=wavelengths)
        self.write_spectrum("H2O_reference", np.zeros(70), wavelengths=wavelengths, time="19:00:00")
        for arguments, window in (([], 50), (["6"], 6)):
            with self.subTest(window=window):
                result = subprocess.run(
                    [sys.executable, str(PROJECT_ROOT / "merge.py"),
                     "--folder", str(self.folder), "--smooth", *arguments,
                     "--subtract", "H2O"],
                    capture_output=True, text=True, check=False,
                )
                self.assertEqual(result.returncode, 0, result.stderr)
                workbook = load_workbook(
                    self.folder.parent / "instrument data-combined.xlsx", data_only=False
                )
                try:
                    actual = [workbook["CD_subtr"].cell(row, 2).value for row in range(2, 72)]
                    np.testing.assert_allclose(actual, quadratic_window_fit(values, window), atol=1e-12)
                    self.assertEqual(workbook["CD_subtr"]["E2"].value, "=B2-$C2")
                finally:
                    workbook.close()

    def test_invalid_smoothing_windows_do_not_write_outputs(self):
        for window in (-1, 0, 1, 2, 4, 1_000_000):
            with self.subTest(window=window):
                with self.assertRaisesRegex(ValueError, "--smooth window"):
                    merge.merge_spectra(self.folder, smooth=window)
                self.assertFalse((self.folder / "concentration_table.txt").exists())
                self.assertFalse((self.folder.parent / "instrument data-combined.xlsx").exists())

    def test_smoothing_rejects_non_numeric_or_non_finite_cd(self):
        for value, expected in (("bad", "numeric"), (np.nan, "finite"), (np.inf, "finite")):
            with self.subTest(value=value):
                with self.assertRaisesRegex(ValueError, f"invalid.csv must be {expected}"):
                    merge.smooth_cd_spectra({"invalid": pd.Series([1, value, 3])}, 3)

    def test_ht_threshold_blanks_only_above_limit_for_each_spectrum(self):
        self.write_spectrum("sample", [5, 8, 11], ht=[499, 500, 501])
        self.write_spectrum("later sample", [2, 4, 6], ht=[501, 490, 500], time="19:00:00")
        self.write_coefficients("sample\t2\nlater sample\t4\n")
        output = merge.merge_spectra(self.folder, ht_threshold=500)
        workbook = load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)
        for column, expected in ((2, [5, 8, None]), (3, [None, 4, 6])):
            actual = [workbook["CD"].cell(row, column).value for row in range(2, 5)]
            self.assertEqual(actual, expected)
        self.assertEqual([cell.value for cell in workbook["CD"]["A"]], ["wavelength_nm", 200, 201, 202])
        self.assertEqual([workbook["HT"].cell(row, 2).value for row in range(2, 5)], [499, 500, 501])
        self.assertEqual([workbook["Abs"].cell(row, 2).value for row in range(2, 5)], [0.25] * 3)
        sheet = workbook["CD_subtr"]
        self.assertEqual([sheet.cell(row, 5).value for row in range(2, 5)], [2.5, 4, None])
        self.assertEqual([sheet.cell(row, 6).value for row in range(2, 5)], [None, 1, 1.5])

    def test_ht_threshold_excludes_invalid_sample_or_reference_from_formulas_and_plot(self):
        self.write_spectrum("sample", [5, 8, 11], ht=[499, 501, 499])
        self.write_spectrum("H2O_reference", [1, 2, 3], ht=[501, 499, 499], time="19:00:00")
        self.write_spectrum("later sample", [4, 6, 8], ht=[499, 499, 501], time="20:00:00")
        self.write_coefficients("sample\t2\nH2O_reference\t3\nlater sample\t1\n")

        def inspect_chart(*args, **kwargs):
            axis = plt.gcf().axes[0]
            self.assertIn("HT cutoff: 500 V", axis.get_title())
            lines = axis.get_lines()
            self.assertEqual([line.get_label() for line in lines], ["sample", "later sample"])
            np.testing.assert_allclose(lines[0].get_ydata(), [np.nan, np.nan, 4], equal_nan=True)
            np.testing.assert_allclose(lines[1].get_ydata(), [np.nan, 4, np.nan], equal_nan=True)

        with patch.object(plt, "show", side_effect=inspect_chart):
            output = merge.merge_spectra(self.folder, subtract="H2O", ht_threshold=500, plot=True)
        workbook = load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)
        sheet = workbook["CD_subtr"]
        for coordinate in ("F2", "G2", "F3", "G4"):
            self.assertIsNone(sheet[coordinate].value)
            self.assertNotEqual(sheet[coordinate].data_type, "f")
        self.assertEqual(sheet["F4"].value, "=(B4-$C4)/2.0")
        self.assertEqual(sheet["G3"].value, "=D3-$C3")
        self.assertEqual(sheet["B2"].value, 5)
        self.assertIsNone(sheet["C2"].value)

    def test_ht_filtering_precedes_smoothing_and_keeps_gaps(self):
        positions = np.arange(31, dtype=float)
        wavelengths = positions + 200
        sample = np.sin(positions / 3) + 0.1 * (-1.0)**positions
        sample[[0, 10, 30]] = 1e9
        ht = np.full(31, 400)
        ht[[0, 10, 30]] = 501
        self.write_spectrum("sample", sample, wavelengths=wavelengths, ht=ht)
        self.write_spectrum("H2O_reference", np.zeros(31), wavelengths=wavelengths, time="19:00:00", ht=300)
        self.write_coefficients("sample\t2\nH2O_reference\t1\n")
        expected = np.full(31, np.nan)
        expected[1:10] = quadratic_window_fit(sample[1:10], 6)
        expected[11:30] = quadratic_window_fit(sample[11:30], 6)

        def inspect_chart(*args, **kwargs):
            lines = plt.gcf().axes[0].get_lines()
            np.testing.assert_allclose(lines[0].get_ydata(), expected / 2, atol=1e-12, equal_nan=True)

        with patch.object(plt, "show", side_effect=inspect_chart):
            output = merge.merge_spectra(self.folder, smooth=6, ht_threshold=500, subtract="H2O", plot=True)
        workbook = load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)
        actual = [workbook["CD_subtr"].cell(row, 2).value for row in range(2, 33)]
        np.testing.assert_allclose(np.array(actual, dtype=float), expected, atol=1e-12, equal_nan=True)
        for position in (0, 10, 30):
            self.assertIsNone(workbook["CD"].cell(position + 2, 2).value)
            self.assertIsNone(workbook["CD_subtr"].cell(position + 2, 5).value)
        self.assertEqual(workbook["CD"].cell(3, 2).value, sample[1])

    def test_smoothing_keeps_short_valid_runs_and_all_excluded_points(self):
        for values in (
            [np.nan, 2, 3, np.nan, 5, 6, np.nan],
            [np.nan] * 7,
        ):
            with self.subTest(values=values):
                result = merge.smooth_cd_spectra({"masked": pd.Series(values)}, 5, allow_gaps=True)["masked"]
                np.testing.assert_allclose(result, values, equal_nan=True)

    def test_all_cd_points_can_be_excluded_with_smoothing(self):
        self.write_spectrum("sample", [5, 8, 11], ht=501)
        output = merge.merge_spectra(self.folder, ht_threshold=500, smooth=3)
        workbook = load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual([workbook["CD"].cell(row, 2).value for row in range(2, 5)], [None] * 3)
        self.assertEqual([workbook["CD_subtr"].cell(row, 4).value for row in range(2, 5)], [None] * 3)

    def test_cli_ht_threshold_supports_default_and_custom_values(self):
        self.write_spectrum("sample", [5, 8, 11], ht=[500, 500.25, 500.75])
        self.write_spectrum("H2O_reference", [1, 1, 1], ht=300, time="19:00:00")
        for arguments, expected in (([], [5, None, None]), (["500.5"], [5, 8, None])):
            with self.subTest(arguments=arguments):
                result = subprocess.run(
                    [sys.executable, str(PROJECT_ROOT / "merge.py"),
                     "--folder", str(self.folder), "--ht-threshold", *arguments,
                     "--subtract", "H2O"],
                    capture_output=True, text=True, check=False,
                )
                self.assertEqual(result.returncode, 0, result.stderr)
                workbook = load_workbook(self.folder.parent / "instrument data-combined.xlsx", data_only=False)
                try:
                    self.assertEqual([workbook["CD"].cell(row, 2).value for row in range(2, 5)], expected)
                    self.assertEqual(workbook["CD_subtr"]["E2"].value, "=B2-$C2")
                    self.assertIsNone(workbook["CD_subtr"]["E4"].value)
                finally:
                    workbook.close()

    def test_invalid_ht_thresholds_do_not_write_outputs(self):
        for threshold in (-1, np.nan, np.inf, -np.inf):
            with self.subTest(threshold=threshold):
                with self.assertRaisesRegex(ValueError, "--ht-threshold value must be finite"):
                    merge.merge_spectra(self.folder, ht_threshold=threshold)
                self.assertFalse((self.folder / "concentration_table.txt").exists())
                self.assertFalse((self.folder.parent / "instrument data-combined.xlsx").exists())

    def test_ht_filtering_rejects_invalid_ht_values(self):
        for value, expected in (("bad", "numeric"), (np.nan, "finite"), (np.inf, "finite")):
            with self.subTest(value=value):
                self.write_spectrum("sample", [5, 8, 11], ht=[400, value, 500])
                with self.assertRaisesRegex(ValueError, f"sample.csv must be {expected}"):
                    merge.merge_spectra(self.folder, ht_threshold=500)

    def test_invalid_concentration_coefficients_are_rejected(self):
        for coefficient in ("0", "-1", "NaN", "inf", "-inf"):
            with self.subTest(coefficient=coefficient):
                self.write_coefficients(f"sample\t{coefficient}\n")
                with self.assertRaisesRegex(ValueError, "finite and greater than zero"):
                    merge.load_concentration_coefficients(self.folder, ["sample"])

    def test_concentration_table_must_match_samples(self):
        for rows, error in (
            ("sample\t1\nsample\t2\n", "Duplicate spectra"),
            ("other\t1\n", "not present in data folder: other"),
        ):
            with self.subTest(rows=rows):
                self.write_coefficients(rows)
                with self.assertRaisesRegex(ValueError, error):
                    merge.load_concentration_coefficients(self.folder, ["sample"])

    def test_empty_selection_and_invalid_short_names_leave_existing_outputs_untouched(self):
        output = merge.merge_spectra(self.folder)
        original_workbook = output.read_bytes()
        table_file = self.folder / "concentration_table.txt"
        self.write_spectrum("second", [4, 6, 8])
        for rows, error in (
            ("", "No sample spectra listed"),
            ("sample\t1\t7\tduplicate\nsecond\t2\t7\tduplicate\n", "Duplicate short names"),
            ("sample\t1\t7\tzero\n", "conflict with workbook headers"),
            ("sample\t1\t7\twavelength_nm\n", "conflict with workbook headers"),
            ("missing\t1\t7\tMissing\n", "not present in data folder: missing"),
        ):
            with self.subTest(rows=rows):
                contents = "\t".join(merge.CONCENTRATION_COLUMNS) + "\n" + rows
                table_file.write_text(contents, encoding="utf-8")
                original_table = table_file.read_bytes()
                with self.assertRaisesRegex(ValueError, error):
                    merge.merge_spectra(self.folder)
                self.assertEqual(table_file.read_bytes(), original_table)
                self.assertEqual(output.read_bytes(), original_workbook)

    def test_missing_markers_and_empty_blocks_are_rejected(self):
        file = self.folder / "malformed.csv"
        for content, error in (
            ("200,1,500\n##### Extended Information\n", "XYDATA marker"),
            ("XYDATA\n200,1,500\n", "Extended Information marker"),
            ("XYDATA\n\n##### Extended Information\n", "No spectral data"),
        ):
            with self.subTest(content=content):
                file.write_text(content, encoding="latin-1")
                with self.assertRaisesRegex(ValueError, error):
                    merge.read_spectrum(file)

    def test_cli_accepts_windows_folder_with_spaces_and_trailing_separator(self):
        result = subprocess.run(
            [sys.executable, str(PROJECT_ROOT / "merge.py"),
             "--folder", str(self.folder) + "\\"],
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Saved ", result.stdout)

    def test_cli_reports_missing_folder_without_traceback(self):
        result = subprocess.run(
            [sys.executable, str(PROJECT_ROOT / "merge.py"),
             "--folder", str(self.folder / "missing")],
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(result.returncode, 1)
        self.assertIn("ERROR: Data folder not found", result.stderr)
        self.assertNotIn("Traceback", result.stderr)


class AnalysisTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="ECD analysis ")
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name)
        wavelengths = np.linspace(190, 260, 20)
        components = np.vstack([
            np.exp(-((wavelengths - 210) / 12) ** 2),
            np.exp(-((wavelengths - 240) / 10) ** 2),
        ])
        spectra = np.array([[1.0, 0.2], [0.4, 1.0], [0.8, 0.8]]) @ components
        self.frame = pd.DataFrame(spectra.T, columns=["sample 1", "sample 2", "sample 3"])
        self.frame.insert(0, "wavelength_nm", wavelengths)
        self.workbook = self.folder / "test spectra.xlsx"
        self.frame.to_excel(self.workbook, index=False, sheet_name="Sheet1")

    def test_outside_range_reports_observed_bounds_with_one_workbook_read(self):
        with patch.object(analysis.pd, "read_excel", wraps=pd.read_excel) as read_excel:
            with self.assertRaisesRegex(ValueError, "190-260 nm"):
                analysis.load_spectral_dataset(self.workbook, "Sheet1", 300, 400)
        self.assertEqual(read_excel.call_count, 1)

    def test_cli_exports_headless_pca_and_mcr_results(self):
        self.run_analysis("positive", expected_csv=8, expected_png=9)

    def test_cli_skips_mcr_for_signed_ecd_data(self):
        self.frame.iloc[:, 1:] -= 0.5
        self.frame.to_excel(self.workbook, index=False, sheet_name="Sheet1")
        result = self.run_analysis("signed", expected_csv=5, expected_png=6)
        self.assertIn("MCR-ALS was skipped", result.stderr)

    def run_analysis(self, name, expected_csv, expected_png):
        output = self.folder / name
        result = subprocess.run(
            [sys.executable, str(PROJECT_ROOT / "spectral_component_analysis.py"),
             "--file", str(self.workbook), "--components", "2",
             "--output", str(output), "--no-show"],
            env={**os.environ, "MPLBACKEND": "Agg"},
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(len(list(output.glob("*.csv"))), expected_csv)
        self.assertEqual(len(list(output.glob("*.png"))), expected_png)
        reconstructed = pd.read_csv(output / "pca_reconstructed_spectra.csv")
        np.testing.assert_allclose(reconstructed, self.frame, atol=1e-12)
        return result


if __name__ == "__main__":
    unittest.main()
